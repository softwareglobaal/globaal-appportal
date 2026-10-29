"""Controle op bijlagen lezen en downloaden: wat eruit komt, en wat nooit kan.

Draaien kan zonder pytest en zonder mailserver:

    python post/test_bijlagen.py

De IMAP-kant wordt vervangen door een vast bericht in het geheugen; er wordt
niets verbonden en er gaat niets uit.
"""
import ast
import base64
import email
import io
import json
import os
import sys
import tempfile
import time
import zipfile
from email.message import EmailMessage
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ["MCP_SECRET"] = "test-geheim"

import yaml            # noqa: E402
from PIL import Image  # noqa: E402

import bijlagen        # noqa: E402
import config          # noqa: E402
import imapbron        # noqa: E402

MAP = tempfile.mkdtemp(prefix="postbus-bijlagen-")
config.PAD = os.path.join(MAP, "mailboxen.yaml")
config.TOEGANG_PAD = os.path.join(MAP, "toegang", "toegang.yaml")
config.LOG_PAD = os.path.join(MAP, "toegang", "wijzigingen.jsonl")
with open(config.PAD, "w", encoding="utf-8") as f:
    yaml.safe_dump({
        "standaard": {"imap_host": "imap.one.com"},
        "mailboxen": [
            {"adres": "info@energie-efficient.be", "wachtwoord": "g",
             "personen": ["mukesh", "joan"], "bijlagen": ["mukesh"],
             "beheer": ["mukesh"]},
            {"adres": "office@energie-efficient.be", "wachtwoord": "g",
             "personen": ["mukesh"]},
        ]}, f)


def gelijk(gekregen, verwacht, wat):
    if gekregen != verwacht:
        raise AssertionError(f"{wat}: verwacht {verwacht!r}, "
                             f"gekregen {gekregen!r}")
    print(f"  ok  {wat}")


def weigert(functie, stuk, wat):
    try:
        functie()
    except ValueError as e:
        if stuk.lower() not in str(e).lower():
            raise AssertionError(
                f"{wat}: geweigerd, maar de melding noemt '{stuk}' niet: {e}")
        print(f"  ok  {wat}")
        return
    raise AssertionError(f"{wat}: dit had geweigerd moeten worden")


# ---- voorbeeldbestanden ----------------------------------------------
def tekst_pdf(regel="Factuur 2026-118 totaal 1.250,00 EUR"):
    stroom = f"BT /F1 18 Tf 72 720 Td ({regel}) Tj ET".encode()
    objecten = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stroom)).encode() + b" >>\nstream\n"
        + stroom + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    uit, posities = io.BytesIO(), []
    uit.write(b"%PDF-1.4\n")
    for i, obj in enumerate(objecten, start=1):
        posities.append(uit.tell())
        uit.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = uit.tell()
    uit.write(f"xref\n0 {len(objecten) + 1}\n0000000000 65535 f \n".encode())
    for p in posities:
        uit.write(f"{p:010d} 00000 n \n".encode())
    uit.write(f"trailer\n<< /Size {len(objecten) + 1} /Root 1 0 R >>\n"
              f"startxref\n{xref}\n%%EOF\n".encode())
    return uit.getvalue()


def beeld(fmt="PNG", maat=(800, 600)):
    uit = io.BytesIO()
    Image.new("RGB", maat, (200, 30, 30)).save(uit, fmt)
    return uit.getvalue()


def scan_pdf(paginas=1):
    uit = io.BytesIO()
    beelden = [Image.new("RGB", (1240, 1754), (240, 240, 240))
               for _ in range(paginas)]
    beelden[0].save(uit, "PDF", save_all=True, append_images=beelden[1:])
    return uit.getvalue()


def docx_bytes():
    import docx
    d = docx.Document()
    d.add_paragraph("Verslag EPB")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Wand", "U-waarde"
    t.cell(1, 0).text, t.cell(1, 1).text = "Gevel", "0,24"
    uit = io.BytesIO()
    d.save(uit)
    return uit.getvalue()


def xlsx_bytes():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Offerte"
    wb.active.append(["Post", "Bedrag"])
    wb.active.append(["Isolatie", 4200])
    uit = io.BytesIO()
    wb.save(uit)
    return uit.getvalue()


def zip_bytes():
    uit = io.BytesIO()
    with zipfile.ZipFile(uit, "w") as z:
        z.writestr("plan/grondplan.pdf", b"x" * 2048)
        z.writestr("foto1.jpg", b"y" * 1024)
    return uit.getvalue()


def bericht():
    binnen = EmailMessage()
    binnen["Subject"] = "Doorgestuurde factuur"
    binnen["From"] = "leverancier@voorbeeld.be"
    binnen.set_content("Zie bijlage.")
    binnen.add_attachment(b"%PDF-1.4 binnenin", maintype="application",
                          subtype="pdf", filename="binnen.pdf")
    m = EmailMessage()
    m["Subject"] = "Stavingsstukken"
    m["From"] = "klant@voorbeeld.be"
    m.set_content("In bijlage de stukken.")
    m.add_attachment(tekst_pdf(), maintype="application", subtype="pdf",
                     filename="factuur.pdf")
    m.add_attachment(beeld("JPEG"), maintype="image", subtype="jpeg",
                     filename="gevel.jpg")
    m.add_attachment(b"<script>alert(1)</script>", maintype="text",
                     subtype="html", filename="factuur.pdf.html")
    m.add_attachment(binnen)
    return email.message_from_bytes(m.as_bytes())


# ---- tests -------------------------------------------------------------
def test_nummering_telt_doorgestuurde_mail_als_een():
    delen = imapbron._bijlage_delen(bericht())
    gelijk([d[0] for d in delen],
           ["factuur.pdf", "gevel.jpg", "factuur.pdf.html",
            "Doorgestuurde factuur.eml"],
           "vier bijlagen, de bijgevoegde mail als een .eml")
    genummerd = [(i, *d) for i, d in enumerate(delen, start=1)]
    gelijk(bijlagen.kies(genummerd, "2")[1], "gevel.jpg", "kiezen op nummer")
    gelijk(bijlagen.kies(genummerd, "gevel")[1], "gevel.jpg",
           "kiezen op stuk naam")
    gelijk(bijlagen.kies(genummerd, "factuur.pdf")[0], 1,
           "exacte naam wint van een langere")
    weigert(lambda: bijlagen.kies(genummerd, "9"), "bestaat niet",
            "onbestaand nummer")


def test_pdf_met_tekst():
    uit = bijlagen.inlezen("factuur.pdf", "application/pdf", tekst_pdf())
    gelijk("1.250,00 EUR" in uit["tekst"], True, "tekst uit de pdf")
    gelijk((uit["scanpaginas"], uit["afbeeldingen"]), (0, []),
           "geen afbeelding bij een tekst-pdf")


def test_pdf_scan_wordt_beeld_in_delen():
    uit = bijlagen.inlezen("scan.pdf", "application/pdf", scan_pdf(7))
    gelijk((uit["paginas"], uit["scanpaginas"], uit["aantal_delen"]),
           (7, 7, 2), "zeven scanpagina's, twee delen")
    gelijk(len(uit["afbeeldingen"]), 5, "deel 1 heeft vijf beelden")
    b = Image.open(io.BytesIO(uit["afbeeldingen"][0][1]))
    gelijk(max(b.size) <= bijlagen.LANGE_ZIJDE, True, "beeld verkleind")
    tweede = bijlagen.inlezen("scan.pdf", "application/pdf", scan_pdf(7), 2)
    gelijk(len(tweede["afbeeldingen"]), 2, "deel 2 heeft de rest")
    weigert(lambda: bijlagen.inlezen("scan.pdf", "application/pdf",
                                     scan_pdf(7), 3),
            "bestaat niet", "deel 3 bestaat niet")


def test_afbeeldingen():
    uit = bijlagen.inlezen("foto.png", "image/png", beeld("PNG", (3000, 2000)))
    b = Image.open(io.BytesIO(uit["afbeeldingen"][0][1]))
    gelijk((b.format, max(b.size)), ("JPEG", bijlagen.LANGE_ZIJDE),
           "png wordt verkleinde jpeg")
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        heic = beeld("HEIF")
    except Exception as e:
        print(f"  --  heic overgeslagen: kan hier geen heic maken ({e})")
    else:
        uit = bijlagen.inlezen("IMG_0001.HEIC", "image/heic", heic)
        gelijk(uit["afbeeldingen"][0][0], "image/jpeg", "heic wordt jpeg")
    oud = bijlagen.MAX_PIXELS
    bijlagen.MAX_PIXELS = 1000
    try:
        weigert(lambda: bijlagen.inlezen("groot.png", "image/png", beeld()),
                "te groot", "te veel pixels wordt geweigerd")
    finally:
        bijlagen.MAX_PIXELS = oud


def test_kantoorbestanden():
    uit = bijlagen.inlezen("verslag.docx", "application/octet-stream",
                           docx_bytes())
    gelijk(("Verslag EPB" in uit["tekst"], "Gevel | 0,24" in uit["tekst"]),
           (True, True), "docx: tekst en tabel")
    uit = bijlagen.inlezen("offerte.xlsx", "application/octet-stream",
                           xlsx_bytes())
    gelijk(("=== blad: Offerte ===" in uit["tekst"],
            "Isolatie | 4200" in uit["tekst"]), (True, True), "xlsx: blad en rij")
    uit = bijlagen.inlezen("lijst.csv", "text/csv", "a;b\n1;é\n".encode("cp1252"))
    gelijk(uit["tekst"], "a;b\n1;é", "csv in cp1252 wordt goed gelezen")
    oud = bijlagen.MAX_UITGEPAKT
    bijlagen.MAX_UITGEPAKT = 100
    try:
        weigert(lambda: bijlagen.inlezen("bom.docx", "", docx_bytes()),
                "te groot", "docx die te groot uitpakt wordt geweigerd")
    finally:
        bijlagen.MAX_UITGEPAKT = oud


def test_zip_en_eml_geven_een_lijst():
    uit = bijlagen.inlezen("stukken.zip", "application/zip", zip_bytes())
    gelijk(("plan/grondplan.pdf" in uit["tekst"],
            "niets is uitgepakt" in uit["tekst"]), (True, True), "zip-lijst")
    delen = imapbron._bijlage_delen(bericht())
    uit = bijlagen.inlezen(delen[3][0], delen[3][1], delen[3][2])
    gelijk(("binnen.pdf" in uit["tekst"],
            "Doorgestuurde factuur" in uit["tekst"]), (True, True),
           "eml: koppen en eigen bijlagen")


def test_lange_tekst_in_delen():
    lang = ("regel\n" * 20000).encode()
    uit = bijlagen.inlezen("log.txt", "text/plain", lang)
    gelijk(uit["aantal_delen"] > 1, True, "lange tekst krijgt delen")
    gelijk(uit["tekens_totaal"], len(lang.decode().strip()), "totaal klopt")


def test_onbekend_en_te_groot():
    weigert(lambda: bijlagen.inlezen("setup.exe", "application/x-msdownload",
                                     b"MZ"), "niet ingelezen",
            "exe wordt niet ingelezen")
    oud = bijlagen.MAX_BIJLAGE
    bijlagen.MAX_BIJLAGE = 10
    try:
        weigert(lambda: bijlagen.inlezen("a.txt", "text/plain", b"x" * 11),
                "bijlage_link", "te grote bijlage verwijst naar downloaden")
    finally:
        bijlagen.MAX_BIJLAGE = oud


def test_recht_bijlagen():
    config._vergeet()
    mukesh = {"gebruiker": "mukesh", "groepen": ["postbus"]}
    joan = {"gebruiker": "joan", "groepen": ["postbus"]}
    info = config.zoek("info@energie-efficient.be", mukesh)
    gelijk(config.mag(info, "bijlagen", mukesh), True, "mukesh mag op info@")
    weigert(lambda: config.vereis_bijlagen(info, joan), "alleen voor",
            "joan niet op info@")
    office = config.zoek("office@energie-efficient.be", mukesh)
    weigert(lambda: config.vereis_bijlagen(office, mukesh), "staat daar niet",
            "niemand op office@")
    gelijk("bijlagen" in config.BEWERKBAAR, False,
           "bijlagen kan niet via /toegang")
    weigert(lambda: config.normaliseer({"personen": "mukesh",
                                        "bijlagen": ["mukesh"]}),
            "niet wijzigen", "/toegang weigert het veld bijlagen")


def _met_nepbericht():
    def nep(mailbox, mapnaam, uid):
        bron = bericht()
        return ({"onderwerp": "Stavingsstukken", "van": "klant@voorbeeld.be"},
                [(i, *d) for i, d in enumerate(
                    imapbron._bijlage_delen(bron), start=1)])
    imapbron.bijlagen = nep


def test_mcp_geeft_beelden_en_link():
    _met_nepbericht()
    import app as webapp
    import mcp_server
    c = webapp.app.test_client()
    token = mcp_server._teken({"t": "acc", "u": "mukesh", "g": ["postbus"],
                               "exp": time.time() + 600})

    def roep(naam, args):
        r = c.post("/mcp", json={"jsonrpc": "2.0", "id": 1,
                                 "method": "tools/call",
                                 "params": {"name": naam, "arguments": args}},
                   headers={"Authorization": f"Bearer {token}"})
        return r.get_json()["result"]

    lijst = roep("bijlage", {"mailbox": "info@energie-efficient.be", "uid": 7})
    gelijk(len(json.loads(lijst["content"][0]["text"])["bijlagen"]), 4,
           "zonder keuze: de lijst")
    foto = roep("bijlage", {"mailbox": "info@energie-efficient.be", "uid": 7,
                            "bijlage": "2"})
    gelijk((foto["isError"], foto["content"][1]["type"],
            foto["content"][1]["mimeType"]), (False, "image", "image/jpeg"),
           "afbeelding komt als image-blok")
    base64.b64decode(foto["content"][1]["data"])
    geen = roep("bijlage", {"mailbox": "office@energie-efficient.be",
                            "uid": 7, "bijlage": "1"})
    gelijk(geen["isError"], True, "zonder recht geweigerd")
    link = json.loads(roep("bijlage_link", {
        "mailbox": "info@energie-efficient.be", "uid": 7})["content"][0]["text"])
    gelijk(link["link"].startswith("https://post."), True, "link gemaakt")
    return link["link"].split("/download/", 1)[1]


def test_download():
    dl = test_mcp_geeft_beelden_en_link()
    import app as webapp
    import mcp_server
    c = webapp.app.test_client()
    mukesh = {"X-Authentik-Username": "mukesh",
              "X-Authentik-Groups": "postbus"}
    r = c.get(f"/download/{dl}", headers=mukesh)
    gelijk((r.status_code, b"gevel.jpg" in r.data), (200, True),
           "lijstpagina voor mukesh")
    r = c.get(f"/download/{dl}/3", headers=mukesh)
    gelijk((r.status_code, r.mimetype, r.headers.get("X-Content-Type-Options"),
            "attachment" in r.headers.get("Content-Disposition", ""),
            "factuur.pdf.html" in r.headers.get("Content-Disposition", "")),
           (200, "application/octet-stream", "nosniff", True, True),
           "html-bijlage gaat als download met eigen naam")
    gelijk(c.get(f"/download/{dl}/1", headers={
        "X-Authentik-Username": "joan",
        "X-Authentik-Groups": "postbus"}).status_code, 403,
        "andere gebruiker krijgt 403")
    gelijk(c.get(f"/download/{dl}x/1", headers=mukesh).status_code, 403,
           "geknoeide link krijgt 403")
    oud = mcp_server._teken({"t": "dl", "u": "mukesh",
                             "m": "info@energie-efficient.be", "f": "INBOX",
                             "i": 7, "b": None, "exp": time.time() - 1})
    gelijk(c.get(f"/download/{oud}/1", headers=mukesh).status_code, 403,
           "verlopen link krijgt 403")
    een = mcp_server._teken({"t": "dl", "u": "mukesh",
                             "m": "info@energie-efficient.be", "f": "INBOX",
                             "i": 7, "b": 2, "exp": time.time() + 60})
    gelijk(c.get(f"/download/{een}/1", headers=mukesh).status_code, 404,
           "link voor een bijlage geeft geen andere")
    gelijk(c.get(f"/download/{een}", headers=mukesh).status_code, 200,
           "link voor een bijlage geeft die ene")


def test_bijlagen_module_verandert_niets():
    """bijlagen.py leest alleen: geen IMAP-wijziging, geen proces, geen schijf."""
    bron = (Path(__file__).resolve().parent / "bijlagen.py").read_text(
        encoding="utf-8")
    boom = ast.parse(bron)
    namen = {n.id for n in ast.walk(boom) if isinstance(n, ast.Name)}
    namen |= {n.attr for n in ast.walk(boom) if isinstance(n, ast.Attribute)}
    for verboden in ("subprocess", "system", "popen", "exec", "eval",
                     "extractall", "extract", "_selecteer_schrijfbaar",
                     "store", "expunge", "smtplib"):
        if verboden in namen:
            raise AssertionError(f"bijlagen.py gebruikt {verboden!r}")
    # Image.open leest uit het geheugen; het ingebouwde open (naar schijf) niet.
    if any(isinstance(n, ast.Name) and n.id == "open" for n in ast.walk(boom)):
        raise AssertionError("bijlagen.py opent een bestand op schijf")
    print("  ok  bijlagen.py schrijft, wijzigt en start niets")


def main():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for test in tests:
        print(f"\n{test.__name__}")
        test()
    print(f"\n{len(tests)} onderdelen doorlopen, alles in orde.")


if __name__ == "__main__":
    main()
