"""Postbus: de inhoud van een bijlage leesbaar maken voor Claude.

Drie regels:

1. **Alleen lezen.** Het bericht komt binnen via imapbron.bijlagen (readonly
   select + BODY.PEEK); hier wordt alleen nog in het geheugen gelezen. Er
   wordt niets naar schijf geschreven, niets uitgepakt en niets uitgevoerd:
   geen subprocess, geen macro's, formules alleen als opgeslagen waarde.
2. **Grenzen tegen kwaadwillige bestanden.** Een maximumgrootte per bijlage,
   een plafond op het uitgepakte volume van zip-gebaseerde bestanden (docx,
   xlsx) en op het aantal pixels van een afbeelding.
3. **Niets stil afgekapt.** Lange inhoud komt in genummerde delen, met het
   totaal erbij, zoals bij bericht().
"""
import email
import io
import zipfile

import imapbron

MAX_BIJLAGE = 25 * 1024 * 1024        # bytes per bijlage
MAX_UITGEPAKT = 200 * 1024 * 1024     # uitgepakt volume van docx/xlsx
MAX_PIXELS = 50_000_000               # per afbeelding of gerenderde pagina
MAX_PAGINAS = 300                     # pagina's per pdf die we bekijken
LANGE_ZIJDE = 1568                    # px; groter leest Claude toch niet
BEELDEN_PER_DEEL = 5
TEKENS_PER_DEEL = imapbron.TEKENS_PER_DEEL
# Een pdf-pagina met minder tekst dan dit is vrijwel zeker een scan.
SCAN_ONDER = 25
# Wat een scanpagina "kost" bij het verdelen over delen, in tekens.
BEELD_GEWICHT = 8000

AFBEELDING = {"jpg", "jpeg", "png", "gif", "webp", "bmp", "tif", "tiff",
              "heic", "heif"}
TEKST = {"txt", "csv", "tsv", "md", "json", "xml", "log", "ini", "vcf", "ics"}


def overzicht(delen):
    return [{"nr": nr, "naam": naam, "type": soort_mime,
             "grootte_kb": round(len(inhoud) / 1024, 1)}
            for nr, naam, soort_mime, inhoud in delen]


def kies(delen, keuze):
    """Eén bijlage op nummer of naam; een stukje naam mag als het uniek is."""
    if not delen:
        raise ValueError("Dit bericht heeft geen bijlagen")
    tekst = str(keuze if keuze is not None else "").strip()
    if tekst.isdigit():
        nr = int(tekst)
        for d in delen:
            if d[0] == nr:
                return d
        raise ValueError(f"Bijlage {nr} bestaat niet; dit bericht heeft er "
                         f"{len(delen)}")
    laag = tekst.lower()
    exact = [d for d in delen if d[1].lower() == laag]
    if len(exact) == 1:
        return exact[0]
    deels = [d for d in delen if laag and laag in d[1].lower()]
    if len(deels) == 1:
        return deels[0]
    namen = "; ".join(f"{d[0]}: {d[1]}" for d in delen)
    raise ValueError(("Meer dan een bijlage past op" if deels else
                      "Geen bijlage heet") + f" '{tekst}'. Kies een nummer: {namen}")


def soort(naam, mime):
    ext = naam.rsplit(".", 1)[-1].lower() if "." in naam else ""
    mime = (mime or "").lower()
    if ext == "pdf" or mime == "application/pdf":
        return "pdf"
    if ext in AFBEELDING or mime.startswith("image/"):
        return "afbeelding"
    if ext == "docx":
        return "docx"
    if ext in ("xlsx", "xlsm"):
        return "xlsx"
    if ext == "xls":
        return "xls"
    if ext == "zip" or mime in ("application/zip", "application/x-zip-compressed"):
        return "zip"
    if ext == "eml" or mime == "message/rfc822":
        return "eml"
    if ext in ("htm", "html") or mime == "text/html":
        return "html"
    if ext in TEKST or mime.startswith("text/"):
        return "tekst"
    return "anders"


# ---- hulpjes -----------------------------------------------------------
def _jpeg(beeld):
    """PIL-beeld naar JPEG-bytes, verkleind tot LANGE_ZIJDE."""
    from PIL import Image
    if beeld.mode in ("RGBA", "LA", "P"):
        beeld = beeld.convert("RGBA")
        wit = Image.new("RGB", beeld.size, (255, 255, 255))
        wit.paste(beeld, mask=beeld.split()[-1])
        beeld = wit
    elif beeld.mode != "RGB":
        beeld = beeld.convert("RGB")
    beeld.thumbnail((LANGE_ZIJDE, LANGE_ZIJDE))
    uit = io.BytesIO()
    beeld.save(uit, "JPEG", quality=85)
    return uit.getvalue()


def _zip_veilig(inhoud, wat):
    try:
        with zipfile.ZipFile(io.BytesIO(inhoud)) as z:
            totaal = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile:
        raise ValueError(f"Dit {wat}-bestand is beschadigd of geen {wat}")
    if totaal > MAX_UITGEPAKT:
        raise ValueError(f"Dit {wat}-bestand pakt uit tot "
                         f"{totaal // (1024 * 1024)} MB; te groot om veilig "
                         "te lezen. Download het met bijlage_link.")


def _tekst_blokken(tekst):
    tekst = tekst.strip()
    if not tekst:
        return [("", [])]
    return [(tekst[i:i + TEKENS_PER_DEEL], [])
            for i in range(0, len(tekst), TEKENS_PER_DEEL)]


def _decodeer(inhoud):
    for codering in ("utf-8-sig", "cp1252"):
        try:
            return inhoud.decode(codering)
        except UnicodeDecodeError:
            continue
    return inhoud.decode("latin-1")


# ---- per soort -----------------------------------------------------------
def _pdf(inhoud, deel):
    import pypdfium2 as pdfium
    try:
        pdf = pdfium.PdfDocument(inhoud)
    except pdfium.PdfiumError as e:
        raise ValueError(f"Deze pdf kan niet geopend worden (beveiligd met "
                         f"een wachtwoord of beschadigd): {e}")
    try:
        totaal = len(pdf)
        bekeken = min(totaal, MAX_PAGINAS)
        paginas = []          # (nummer, tekst, is_scan)
        for i in range(bekeken):
            pagina = pdf[i]
            tp = pagina.get_textpage()
            tekst = (tp.get_text_bounded() or "").replace("\r\n", "\n").strip()
            tp.close()
            pagina.close()
            paginas.append((i + 1, tekst, len(tekst) < SCAN_ONDER))

        # Pagina's vast over delen verdelen, zodat deel N altijd dezelfde
        # pagina's is: tekst tot TEKENS_PER_DEEL, scans tot BEELDEN_PER_DEEL.
        groepen, huidig, kost, beelden = [], [], 0, 0
        for p in paginas:
            gewicht = BEELD_GEWICHT if p[2] else len(p[1])
            if huidig and (kost + gewicht > TEKENS_PER_DEEL
                           or (p[2] and beelden >= BEELDEN_PER_DEEL)):
                groepen.append(huidig)
                huidig, kost, beelden = [], 0, 0
            huidig.append(p)
            kost += gewicht
            beelden += 1 if p[2] else 0
        if huidig or not groepen:
            groepen.append(huidig)
        if deel > len(groepen):
            raise ValueError(f"Deel {deel} bestaat niet; deze pdf heeft "
                             f"{len(groepen)} deel(en)")

        stukken, afbeeldingen = [], []
        for nummer, tekst, is_scan in groepen[deel - 1]:
            if is_scan:
                pagina = pdf[nummer - 1]
                b, h = pagina.get_size()
                schaal = min(LANGE_ZIJDE / max(b, h, 1), 200 / 72)
                if b * h * schaal * schaal > MAX_PIXELS:
                    schaal = (MAX_PIXELS / (b * h)) ** 0.5
                beeld = pagina.render(scale=schaal).to_pil()
                pagina.close()
                afbeeldingen.append(("image/jpeg", _jpeg(beeld)))
                stukken.append(f"--- pagina {nummer} (scan, zie afbeelding "
                               f"{len(afbeeldingen)}) ---")
            else:
                stukken.append(f"--- pagina {nummer} ---\n{tekst}")
        let_op = None
        if bekeken < totaal:
            let_op = (f"Deze pdf heeft {totaal} pagina's; alleen de eerste "
                      f"{bekeken} zijn ingelezen. Download het bestand voor "
                      "de rest.")
        scans = sum(1 for p in paginas if p[2])
        return {"paginas": totaal, "scanpaginas": scans,
                "aantal_delen": len(groepen), "tekst": "\n\n".join(stukken),
                "afbeeldingen": afbeeldingen, "let_op": let_op}
    finally:
        pdf.close()


def _afbeelding(inhoud):
    from PIL import Image, ImageOps
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except ImportError:
        pass
    try:
        beeld = Image.open(io.BytesIO(inhoud))
    except Exception as e:
        raise ValueError(f"Deze afbeelding kan niet geopend worden: {e}")
    b, h = beeld.size
    if b * h > MAX_PIXELS:
        raise ValueError(f"Deze afbeelding is {b}x{h} pixels; te groot om "
                         "veilig te openen. Download ze met bijlage_link.")
    frames = getattr(beeld, "n_frames", 1)
    beeld = ImageOps.exif_transpose(beeld)
    return {"breedte": b, "hoogte": h, "aantal_delen": 1, "tekst": "",
            "afbeeldingen": [("image/jpeg", _jpeg(beeld))],
            "let_op": (f"Dit bestand heeft {frames} beelden; alleen het "
                       "eerste wordt getoond.") if frames > 1 else None}


def _docx(inhoud):
    _zip_veilig(inhoud, "docx")
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(io.BytesIO(inhoud))
    regels = []
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            regels.append(Paragraph(el, doc).text)
        elif tag == "tbl":
            regels.append("[tabel]")
            for rij in Table(el, doc).rows:
                regels.append(" | ".join(c.text.strip() for c in rij.cells))
            regels.append("[einde tabel]")
    return "\n".join(regels)


def _xlsx(inhoud):
    _zip_veilig(inhoud, "xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(inhoud), read_only=True,
                                data_only=True)
    try:
        regels = []
        for blad in wb.worksheets:
            regels.append(f"=== blad: {blad.title} ===")
            for rij in blad.iter_rows(values_only=True):
                waarden = ["" if v is None else str(v) for v in rij]
                while waarden and waarden[-1] == "":
                    waarden.pop()
                if waarden:
                    regels.append(" | ".join(waarden))
        return "\n".join(regels)
    finally:
        wb.close()


def _xls(inhoud):
    import xlrd
    try:
        wb = xlrd.open_workbook(file_contents=inhoud, on_demand=True)
    except xlrd.XLRDError as e:
        raise ValueError(f"Dit xls-bestand kan niet gelezen worden: {e}")
    regels = []
    for blad in wb.sheets():
        regels.append(f"=== blad: {blad.name} ===")
        for r in range(blad.nrows):
            waarden = []
            for cel in blad.row(r):
                if cel.ctype == xlrd.XL_CELL_DATE:
                    try:
                        waarden.append(str(xlrd.xldate_as_datetime(
                            cel.value, wb.datemode)))
                        continue
                    except Exception:
                        pass
                waarden.append("" if cel.value in (None, "") else str(cel.value))
            while waarden and waarden[-1] == "":
                waarden.pop()
            if waarden:
                regels.append(" | ".join(waarden))
    wb.release_resources()
    return "\n".join(regels)


def _zip(inhoud):
    try:
        z = zipfile.ZipFile(io.BytesIO(inhoud))
    except zipfile.BadZipFile:
        raise ValueError("Deze zip is beschadigd of geen zip")
    with z:
        rijen = [i for i in z.infolist() if not i.is_dir()]
        regels = [f"Zip met {len(rijen)} bestand(en); niets is uitgepakt."]
        for i in rijen:
            slot = " (met wachtwoord)" if i.flag_bits & 0x1 else ""
            regels.append(f"- {i.filename}  {round(i.file_size / 1024, 1)} kB"
                          f"{slot}")
    return "\n".join(regels)


def _eml(inhoud):
    bron = email.message_from_bytes(inhoud)
    tekst, _ = imapbron._tekst_en_bijlagen(bron)
    binnen = imapbron._bijlage_delen(bron)
    regels = [f"Van: {imapbron._kop(bron.get('From'))}",
              f"Aan: {imapbron._kop(bron.get('To'))}",
              f"Datum: {imapbron._kop(bron.get('Date'))}",
              f"Onderwerp: {imapbron._kop(bron.get('Subject'))}",
              "", f"Bijlagen in deze mail ({len(binnen)}):"]
    regels += [f"- {n}  ({t}, {round(len(b) / 1024, 1)} kB)"
               for n, t, b in binnen] or ["- geen"]
    regels += ["", "Tekst:", tekst]
    return "\n".join(regels)


def inlezen(naam, mime, inhoud, deel=1):
    """De inhoud van een bijlage: tekst en/of afbeeldingen, in delen."""
    deel = max(1, int(deel or 1))
    if len(inhoud) > MAX_BIJLAGE:
        raise ValueError(f"Deze bijlage is {len(inhoud) // (1024 * 1024)} MB; "
                         f"lezen gaat tot {MAX_BIJLAGE // (1024 * 1024)} MB. "
                         "Download ze met bijlage_link.")
    s = soort(naam, mime)
    if s == "pdf":
        uit = _pdf(inhoud, deel)
    elif s == "afbeelding":
        if deel > 1:
            raise ValueError("Een afbeelding heeft maar een deel")
        uit = _afbeelding(inhoud)
    elif s == "anders":
        raise ValueError(f"Bestandstype van '{naam}' ({mime}) wordt niet "
                         "ingelezen. Download het met bijlage_link.")
    else:
        lezer = {"docx": _docx, "xlsx": _xlsx, "xls": _xls, "zip": _zip,
                 "eml": _eml, "tekst": _decodeer,
                 "html": lambda b: imapbron._html_naar_tekst(_decodeer(b))}[s]
        try:
            volledig = lezer(inhoud)
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"'{naam}' kan niet gelezen worden: "
                             f"{type(e).__name__}: {e}")
        blokken = _tekst_blokken(volledig)
        if deel > len(blokken):
            raise ValueError(f"Deel {deel} bestaat niet; deze bijlage heeft "
                             f"{len(blokken)} deel(en)")
        uit = {"tekens_totaal": len(volledig.strip()),
               "aantal_delen": len(blokken), "tekst": blokken[deel - 1][0],
               "afbeeldingen": [], "let_op": None}
    uit["soort"] = s
    uit["deel"] = deel
    return uit
