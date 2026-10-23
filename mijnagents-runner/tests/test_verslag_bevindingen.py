"""Grendel: een verslag van bevindingen is geen werfverslag.

Gebeurd met dossier 2309 (Stockemstraat 15, Huldenberg): het bezoek van 3 juni 2026 voor gebreken na werken van een
derde (Roofconcept) kwam op het bord als 'werfbezoek 1' en de schrijver maakte acht keer '2309-1 werfverslag (concept)'
met veiligheidstekst, volgend werfbezoek en akkoord binnen vijf dagen. Mehdi, 21-10-2026: "is niet een WV maar een
verslag van bevinding". Deze toetsen falen zodra zo'n bezoek weer een werfverslagnummer krijgt, naar de werfapp gaat
of de opmaak van een werfverslag krijgt.
"""
import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

HIER = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HIER))
sys.path.insert(0, str(HIER / "koppelingen"))
import sjabloon_bevindingen as vb  # noqa: E402
import sjabloon_werfverslag as sj  # noqa: E402
import werfapp  # noqa: E402
import werfregister  # noqa: E402
import werfregistratie  # noqa: E402
import werfverslag_voorbereider as V  # noqa: E402

MAP = "/Work All/01. H-A WORK/0 H-A Light projects/5. H-A light SITE VISITS/9999 Proefstraat 1/3. Implementation & Site Control/5. Site Reports"


def docx_tekst(b):
    from docx import Document
    d = Document(io.BytesIO(b))
    tekst = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    for s in d.sections:
        tekst += [p.text for p in s.footer.paragraphs] + [c.text for t in s.footer.tables for r in t.rows for c in r.cells]
    return "\n".join(tekst)


def proef_verslag():
    """Zoals de schrijver het aanlevert: een vast werfpunt en twee eigen vaststellingen."""
    return {"dossier": "9999", "bezoek": "VB1", "soort_bezoek": "bevindingen", "adres": "Proefstraat 1, 3000 Leuven",
            "datum": "2026-06-03", "opgemaakt": "2026-10-21", "verslagtype": vb.SOORT, "status_nl": "Bezoek na werken van een derde.",
            "bronnen_kort": "1 document(en), 2 foto's en 1 opname(s) in de bezoekmap",
            "categorieen": [{"naam": "Algemeen", "punten": [{"titel": "Werfbezoeken", "doorlopend": True, "fotos": [], "tekst_nl": "vast"}]},
                            {"naam": "Dakkapel", "punten": [
                                {"titel": "Zinken dorpel oneffen", "tekst_nl": "Water blijft staan.", "bron": "1027 a.jpg", "fotos": ["1027 a.jpg"]},
                                {"titel": "Naad aan het raam", "tekst_nl": "Open naad.", "bron": "", "fotos": []}]}],
            "raming": [{"onderdeel": "Dakkapel", "herstelling": "dorpel vernieuwen", "raming": "(raming door de architect in te vullen)"}]}


class Sjabloon(unittest.TestCase):
    def test_leeg_sjabloon_is_geen_werfverslag(self):
        tekst = docx_tekst(vb.naar_docx(vb.leeg("9999", "Proefstraat 1")))
        self.assertIn("Verslag van bevindingen", tekst)
        self.assertIn("9999 | Verslag van bevindingen", tekst)
        for werfverslagtekst in ("Werfverslag", sj.VAST_VEILIGHEID[:40], "Peeters-Borsus", "vijf kalenderdagen", "Volgend werfbezoek"):
            self.assertNotIn(werfverslagtekst, tekst)

    def test_herkomst_en_raming(self):
        self.assertEqual(vb._herkomst("V + B"), "eigen vaststelling ter plaatse, verklaring van de bouwheer")
        self.assertEqual(vb._herkomst("transcript.txt 00:22:30"), "transcript.txt 00:22:30")
        self.assertEqual(vb._herkomst(""), "(na te kijken)")
        v = {"raming": [{"van": 100, "tot": 200}, {"van": 7000, "tot": 9000, "variant": True}, {"tekst": "open"}]}
        self.assertEqual(vb.raming_totaal(v), (100, 200))

    def test_proef_van_de_schrijver_wordt_verslag_van_bevindingen(self):
        v = vb.uit_werfproef(proef_verslag())
        self.assertEqual(v["referentie"], "9999-VB1")
        nrs = [b["nr"] for o in v["onderdelen"] for b in o["bevindingen"]]
        self.assertEqual(nrs, ["B1", "B2"])                     # het vaste werfpunt valt weg
        self.assertEqual(v["onderdelen"][0]["bevindingen"][1]["herkomst"], "(na te kijken)")
        self.assertEqual(vb.als_werfverslag(v)["categorieen"][0]["punten"][0]["fotos"], ["1027 a.jpg"])
        md = vb.naar_markdown(v)
        self.assertIn("(raming door de architect in te vullen)", md)
        self.assertGreater(sj.telling_open(md)["in_te_vullen"], 0)


class Nummering(unittest.TestCase):
    def test_bevindingen_telt_niet_als_werfbezoek(self):
        bezoeken = [{"map": MAP + "/20260603", "datum": "2026-06-03"}, {"map": MAP + "/20260910", "datum": "2026-09-10"}]
        rijen = [{"bezoekmap": MAP + "/20260603", "keuzes": {"verslagtype": "verslag_bevindingen"}},
                 {"bezoekmap": MAP + "/20260910", "keuzes": {"verslagtype": "werfverslag"}}]
        V.nummer_bezoeken(bezoeken, "", V.bevindingenmappen(rijen))
        self.assertEqual([(b["soort_bezoek"], b["nr_label"], b["werfnr"]) for b in bezoeken],
                         [("bevindingen", "VB1", 0), ("werfbezoek", "1", 1)])

    def test_zonder_keuze_blijft_de_nummering_zoals_ze_was(self):
        bezoeken = [{"map": MAP + "/20260603", "datum": "2026-06-03"}]
        V.nummer_bezoeken(bezoeken, "")
        self.assertEqual(bezoeken[0]["nr_label"], "1")

    def test_werfapp_weigert_een_vb_nummer(self):
        self.assertIsNone(werfapp.volgnummer("VB1"))

    def test_registratie_vraagt_geen_werfbezoeknummer_voor_bevindingen(self):
        geschreven = []

        def api(pad, payload=None):
            if pad.startswith("/bezoeken?"):
                return {"bezoeken": [], "meer": False}
            geschreven.append(payload)
            return {"bezoek": {"bezoek_id": "vb-1", "project_ref": "h-architects:9999", "status": "onzeker", "revisie": 1}}

        bezoeken = [{"map": MAP + "/20260603", "datum": "2026-06-03"}]
        rijen = [{"bezoekmap": MAP + "/20260603", "keuzes": {"verslagtype": "verslag_bevindingen"}}]
        with patch.object(werfregister, "call", side_effect=api):
            werfregistratie.koppel_voor_nummering(bezoeken, "9999", MAP, "", [], bevindingen=V.bevindingenmappen(rijen))
        self.assertEqual(geschreven[0]["bezoektype"], "overig")
        self.assertEqual(bezoeken[0]["registratie"]["bezoek_id"], "vb-1")

    # driehoeksoverleg 22-10-2026 (K8, K9, C5, G6, G8): verhuis, vast bezoek-ID, twee VB's
    def test_keuze_overleeft_verhuis_van_de_projectmap_met_spaties(self):
        oud, nieuw = "/Work All/Fase 4 wachten/9999 Proef straat 1", "/Work All/Fase 5 loopt/9999 Proef straat 1"
        rel = "/_00. Communication/2026-06-03 bezoek klant"
        rijen = [{"projectmap": oud, "bezoekmap": oud + rel, "keuzes": {"verslagtype": "verslag_bevindingen"}}]
        bezoeken = [{"map": nieuw + rel, "datum": "2026-06-03"}]
        V.nummer_bezoeken(bezoeken, "", V.bevindingenmappen(rijen), nieuw)
        self.assertEqual(bezoeken[0]["nr_label"], "VB1")

    def test_keuze_volgt_het_vaste_bezoek_id(self):
        rijen = [{"bezoekmap": "/elders/oude map", "keuzes": {"verslagtype": "verslag_bevindingen"},
                  "registratie": {"bezoek_id": "vast-7"}}]
        bezoeken = [{"map": MAP + "/hernoemd", "datum": "2026-06-03", "registratie": {"bezoek_id": "vast-7"}}]
        V.nummer_bezoeken(bezoeken, "", V.bevindingenmappen(rijen), MAP)
        self.assertEqual(bezoeken[0]["soort_bezoek"], "bevindingen")

    def test_twee_verslagen_van_bevindingen_tellen_door(self):
        rijen = [{"bezoekmap": MAP + "/a", "keuzes": {"verslagtype": "verslag_bevindingen"}},
                 {"bezoekmap": MAP + "/c", "keuzes": {"verslagtype": "verslag_bevindingen"}}]
        bezoeken = [{"map": MAP + "/a", "datum": "2026-06-03"}, {"map": MAP + "/b", "datum": "2026-07-01"},
                    {"map": MAP + "/c", "datum": "2026-08-01"}]
        V.nummer_bezoeken(bezoeken, "", V.bevindingenmappen(rijen), MAP)
        self.assertEqual([b["nr_label"] for b in bezoeken], ["VB1", "1", "VB2"])
        v = vb.uit_werfproef(dict(proef_verslag(), bezoek="VB2"))
        self.assertEqual(v["referentie"], "9999-VB2")
        self.assertIn("na te kijken", vb.uit_werfproef(dict(proef_verslag(), bezoek="3"))["referentie"])

    def test_registratie_herkent_bevindingen_via_bestaande_bordregistratie(self):
        """Een bezoek dat al als werfbezoek geregistreerd was, houdt zijn ID maar telt niet meer als werfbezoek 1."""
        rij = {"datum": "2026-06-03", "bezoekmap": MAP + "/20260603", "keuzes": {"verslagtype": "verslag_bevindingen"},
               "registratie": {"bezoek_id": "oud-1", "project_ref": "h-architects:9999"}}
        bezoeken = [{"map": MAP + "/20260603", "datum": "2026-06-03"}, {"map": MAP + "/20260910", "datum": "2026-09-10"}]
        aangemaakt = []

        def api(pad, payload=None):
            if pad.startswith("/bezoeken?"):
                return {"bezoeken": [], "meer": False}
            aangemaakt.append(payload)
            return {"bezoek": {"bezoek_id": "nieuw-2", "project_ref": "h-architects:9999"}}

        with patch.object(werfregister, "call", side_effect=api):
            werfregistratie.koppel_voor_nummering(bezoeken, "9999", MAP, "", [rij], bevindingen=V.bevindingenmappen([rij]))
        self.assertEqual(bezoeken[0]["registratie"]["bezoek_id"], "oud-1")
        self.assertEqual(aangemaakt[0]["bezoektype"], "werfbezoek")   # het tweede bezoek is werfbezoek 1


class Fotos(unittest.TestCase):
    def test_iphone_foto_staat_rechtop_in_word(self):
        """PDF van 2309 (23-10-2026): een rechtopstaande foto met EXIF-oriëntatie 6 stond 90 graden gedraaid."""
        from PIL import Image
        buf = io.BytesIO()
        exif = Image.Exif()
        exif[0x0112] = 6
        Image.new("RGB", (40, 30), "white").save(buf, format="JPEG", exif=exif)
        recht = Image.open(io.BytesIO(sj.beeld_rechtop(buf.getvalue())))
        self.assertEqual(recht.size, (30, 40))
        self.assertIn(recht.getexif().get(0x0112, 1), (1, None))
        zonder = io.BytesIO()
        Image.new("RGB", (40, 30), "white").save(zonder, format="JPEG")
        self.assertEqual(sj.beeld_rechtop(zonder.getvalue()), zonder.getvalue())
        self.assertEqual(sj.beeld_rechtop(b"geen beeld"), b"geen beeld")


class Raming(unittest.TestCase):
    def test_twee_totalen_en_pm(self):
        v = {"raming": [{"nr": "R1", "van": 100, "tot": 200}, {"nr": "R3a", "van": 50, "tot": 80},
                        {"nr": "R3b", "van": 700, "tot": 900, "variant": True, "vervangt": "R3a"},
                        {"nr": "R9", "post": "onvoorzien", "tekst": "p.m."}]}
        self.assertEqual(vb.raming_totalen(v), [("Totaal", 150, 280), ("Totaal met R3b in plaats van R3a", 800, 1100)])
        md = vb.naar_markdown(dict(vb.leeg("9999", "x"), raming=v["raming"]))
        self.assertIn("p.m.", md)
        self.assertIn("Totaal met R3b in plaats van R3a", md)


class Schrijver(unittest.TestCase):
    """De grendel kijkt naar de actuele keuze, niet alleen naar de oude proef (C6); --herlayout leest het JSON (C7)."""

    @classmethod
    def setUpClass(cls):
        import werfverslag_proef as P
        cls.P = P

    class Ag:
        def __init__(self):
            self.logs = []

        def log(self, *a):
            self.logs.append(a)

    def rij(self, **extra):
        r = {"id": 1, "dossier": "9999", "datum": "2026-06-03", "bezoekmap": MAP + "/20260603", "adres": "Proefstraat 1",
             "keuzes": {"verslagtype": "verslag_bevindingen"}, "bronnen": {"nr_label": "1", "soort_bezoek": "werfbezoek"},
             "proef_info": {"verslagtype": "vaststellingsverslag", "werfapp": {"id": 41, "nummer": "9999-1"}},
             "verslag_md": "# Werfverslag 1 voor project 9999", "proef_pad": "/x.docx"}
        r.update(extra)
        return r

    def test_werfapp_brug_weigert_bij_keuze_en_toont_oud_concept(self):
        with patch.object(self.P.werfapp, "concept", side_effect=AssertionError("mag niet naar de werfapp")):
            uit = self.P.naar_werfapp(self.Ag(), "9999", 1, self.rij(), {"bezoek": "1"}, [], "/x.docx")
        self.assertIn("overgeslagen", uit)
        self.assertEqual(uit["oud"]["id"], 41)

    def test_doorzetten_naar_werfapp_weigert_bij_keuze(self):
        with patch.object(self.P, "_rij", return_value=self.rij()):
            with self.assertRaisesRegex(RuntimeError, "verslag van bevindingen"):
                self.P.werfapp_uit_proef(self.Ag(), "9999", 1)

    def test_herlayout_leest_het_gegevensbestand_en_overschrijft_niets(self):
        geschreven, bewaard = [], []
        gegevens = vb.leeg("9999", "Proefstraat 1")
        gegevens["referentie"] = "9999-VB1"
        ag = self.Ag()
        with patch.object(self.P, "_rij", return_value=self.rij()), \
             patch.object(self.P.bronnen, "download", return_value=__import__("json").dumps(gegevens).encode()), \
             patch.object(self.P.bronnen, "upload", side_effect=lambda pad, data: geschreven.append(pad) or pad), \
             patch.object(self.P, "lees_bezoekmap", return_value=([], [], [])), \
             patch.object(self.P, "_bewaar", side_effect=lambda *a, **k: bewaard.append(k)):
            pad = self.P.herlayout(ag, "9999", 1)
        self.assertTrue(pad.endswith("9999 Verslag van bevindingen (concept).docx"))
        self.assertEqual(bewaard[0]["proef_info"]["verslagtype"], vb.SOORT)
        self.assertTrue(any("(gegevens).json" in str(x) for x in ag.logs[-1]))


if __name__ == "__main__":
    unittest.main()
