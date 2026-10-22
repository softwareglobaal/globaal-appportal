#!/usr/bin/env python3
"""Gedeelde motor van de verslagagents; werfverslag heeft zijn eigen keten.

Bij barsten en scheuren is de TKN-dossiermap ook de bezoekmap. De wachten vullen
haar bestaande submappen. Een afgeronde levering gaat zonder tokens naar de
inbox van het dashboard. Alleen de knop op het bord start een proef met Claude.
Andere verslagsoorten houden hun bestaande bezoekmap en werkwijze.

Een nieuwe verslagsoort = een blok in koppelingen/verslagsoorten.py + een dunne runner:
    from verslag_basis import VerslagAgent; VerslagAgent("<soort>").main()
"""
import argparse
import fcntl
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urlencode

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HIER)
sys.path.insert(0, os.path.join(HIER, "koppelingen"))
sys.path.insert(0, os.path.join(HIER, "plaud"))
import bord  # noqa: E402
import bronnen  # noqa: E402
import zoektaak as pz  # noqa: E402
import verslagsoorten as vs  # noqa: E402
import werfverslag_proef as wp  # noqa: E402  (leest bezoekmappen, bundelt bronnen, Claude-client, sleutel uit ~/agents/.env)

FOTO_EXT = (".heic", ".jpg", ".jpeg", ".png", ".webp")
OPNAME_EXT = (".mp3", ".m4a", ".wav", ".mp4", ".mov")
DOC_EXT = (".docx", ".pdf", ".pptx", ".md", ".txt")
BS_INBOX = os.environ.get("BS_INBOX_DIR", "/home/ubuntu/bs-inbox")
RUST_MIN = int(os.environ.get("BS_RUST_MIN", "60"))
ZONDER_TAKEN_UUR = 6
MAX_VERSIES = 5
TMP_OUD_UUR = 2
_CONTRACT = vs.pakketcontract()
MAX_BESTAND = _CONTRACT["max_bestand_mb"] * 1_000_000

VERSLAG_SCHEMA = {
    "name": "verslag",
    "description": "Het concept-verslag, per hoofdstuk, met genummerde punten en foto's uit de bezoekmap.",
    "input_schema": {
        "type": "object",
        "properties": {
            "titel": {"type": "string"},
            "aanwezigen": {"type": "array", "items": {"type": "object", "properties": {
                "rol": {"type": "string"}, "firma": {"type": "string"}, "naam": {"type": "string"}, "aanwezig": {"type": "string"}}}},
            "hoofdstukken": {"type": "array", "items": {"type": "object", "properties": {
                "titel": {"type": "string"},
                "tekst": {"type": "string", "description": "lopende tekst van het hoofdstuk; leeg als de punten volstaan"},
                "punten": {"type": "array", "items": {"type": "object", "properties": {
                    "nr": {"type": "string"}, "titel": {"type": "string"}, "tekst": {"type": "string"},
                    "vlag": {"type": "string", "description": "OK, Belangrijk, Dringend of leeg"},
                    "verantwoordelijke": {"type": "string"}, "termijn": {"type": "string"},
                    "fotos": {"type": "array", "items": {"type": "string"}, "description": "bestandsnamen uit de bezoekmap"}},
                    "required": ["titel", "tekst"]}}},
                "required": ["titel"]}},
            "acties": {"type": "array", "items": {"type": "object", "properties": {
                "wie": {"type": "string"}, "wat": {"type": "string"}, "wanneer": {"type": "string"}}}},
            "open_punten": {"type": "array", "items": {"type": "string"},
                            "description": "wat in de bronnen ontbreekt of na te kijken is; Mehdi vult het in"},
        },
        "required": ["titel", "hoofdstukken"],
    },
}


def _systeem(s):
    hoofdstukken = "\n".join(f"{i}. {t}: {oms}" for i, (t, oms) in enumerate(s["hoofdstukken"], 1))
    return f"""Je schrijft voor Mehdi Chegini (architect, {s['afdeling'].upper()}) het concept van een {s['label'].lower()}-verslag
na een bezoek ter plaatse. Je krijgt de bronnen uit de bezoekmap: transcript van de opname, notities, foto's (namen), documenten.

Regels:
- Alleen wat in de bronnen staat. Wat er niet in staat, wordt "(in te vullen)" of komt in open_punten. Verzin niets.
- Elk vaststellingspunt: korte titel, concrete tekst, plaats in het gebouw, en de foto's die erbij horen (bestandsnamen uit de lijst).
- Nummer de punten per hoofdstuk als <hoofdstuk>.<volgnummer> (bv. 3.1, 3.2).
- Onzekere lezing uit het transcript krijgt "(na te kijken)". Sprekers zonder naam heten "spreker N".
- Geen juridische conclusies over aansprakelijkheid; beschrijf vaststellingen en verwijs naar de bevoegde partij.
- Externen noem je met firmanaam; personen alleen als ze in de bronnen met naam en rol staan.
- Nederlands, zakelijk, geen emoji, geen kastlijntjes (gewoon koppelteken).

Hoofdstukken van dit verslagtype, in deze volgorde (allemaal opnemen; een hoofdstuk zonder inhoud krijgt de tekst "geen"):
{hoofdstukken}

Vaste slotzin (niet herschrijven): {s.get('slot', '')}"""


def _md(v, s, rij):
    r = [f"# {v.get('titel') or s.get('verslagtitel', s['label'])}", "",
         f"CONCEPT, opgemaakt {date.today().isoformat()} door {s['agent']} (Mehdi Agents). Mehdi legt er zijn laag over; niets is verstuurd.", "",
         f"- Dossier: {rij.get('_projectnummer') or rij.get('dossier') or '(in te vullen)'}", f"- Adres: {rij.get('adres') or '(in te vullen)'}",
         f"- Klant: {rij.get('klant') or '(in te vullen)'}", f"- Bezoek: {rij.get('datum')} {(rij.get('start') or '')[11:16]}-{(rij.get('einde') or '')[11:16]}",
         f"- Bezoekmap: {rij.get('bezoekmap')}", ""]
    if rij.get("_projectnummer") and str(rij.get("dossier") or "") != rij["_projectnummer"]:
        r.insert(5, f"- Agendanummer: {rij.get('dossier') or '-'}")
    if v.get("aanwezigen"):
        r += ["## Aanwezigen", "", "| Rol | Firma | Naam | Aanwezig |", "|---|---|---|---|"]
        r += [f"| {x.get('rol','')} | {x.get('firma','')} | {x.get('naam','')} | {x.get('aanwezig','')} |" for x in v["aanwezigen"]]
        r.append("")
    for i, h in enumerate(v.get("hoofdstukken") or [], 1):
        r += [f"## {i}. {h.get('titel','')}", ""]
        if h.get("tekst"):
            r += [h["tekst"], ""]
        for pnt in h.get("punten") or []:
            kop = f"**{pnt.get('nr') or ''} {pnt.get('titel','')}**".replace("** ", "**")
            meta = " · ".join(x for x in (pnt.get("vlag") or "", pnt.get("verantwoordelijke") or "", pnt.get("termijn") or "") if x)
            r.append(kop + (f" ({meta})" if meta else ""))
            r.append(pnt.get("tekst", ""))
            if pnt.get("fotos"):
                r.append("Foto's: " + ", ".join(pnt["fotos"]))
            r.append("")
    if v.get("acties"):
        r += ["## Acties", "", "| Wie | Wat | Tegen wanneer |", "|---|---|---|"]
        r += [f"| {x.get('wie','')} | {x.get('wat','')} | {x.get('wanneer','')} |" for x in v["acties"]]
        r.append("")
    if v.get("open_punten"):
        r += ["## Open punten (door Mehdi in te vullen)", ""] + [f"- {x}" for x in v["open_punten"]] + [""]
    if s.get("slot"):
        r += ["---", "", s["slot"], ""]
    return "\n".join(r)


def _docx(v, s, rij, beelden):
    """Eenvoudig Word-document (python-docx, geen master): titel, gegevens, hoofdstukken met punten, foto's erbij."""
    from docx import Document
    from docx.shared import Cm, Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Aptos"
    st.font.size = Pt(11)
    doc.add_heading(v.get("titel") or s.get("verslagtitel", s["label"]), level=0)
    doc.add_paragraph(f"CONCEPT, opgemaakt {date.today().isoformat()}. Dossier {rij.get('_projectnummer') or rij.get('dossier') or '(in te vullen)'} - "
                      f"{rij.get('adres') or ''} - bezoek {rij.get('datum')}")
    if rij.get("_projectnummer") and str(rij.get("dossier") or "") != rij["_projectnummer"]:
        doc.add_paragraph(f"Agendanummer: {rij.get('dossier') or '-'}")
    if v.get("aanwezigen"):
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for c, k in zip(t.rows[0].cells, ("Rol", "Firma", "Naam", "Aanwezig")):
            c.text = k
        for x in v["aanwezigen"]:
            cells = t.add_row().cells
            for c, k in zip(cells, ("rol", "firma", "naam", "aanwezig")):
                c.text = str(x.get(k, ""))
    for i, h in enumerate(v.get("hoofdstukken") or [], 1):
        doc.add_heading(f"{i}. {h.get('titel','')}", level=1)
        if h.get("tekst"):
            doc.add_paragraph(h["tekst"])
        for pnt in h.get("punten") or []:
            p = doc.add_paragraph()
            p.add_run(f"{pnt.get('nr') or ''} {pnt.get('titel','')}".strip()).bold = True
            meta = " · ".join(x for x in (pnt.get("vlag") or "", pnt.get("verantwoordelijke") or "", pnt.get("termijn") or "") if x)
            if meta:
                p.add_run(f"  ({meta})")
            doc.add_paragraph(pnt.get("tekst", ""))
            for naam in pnt.get("fotos") or []:
                b = beelden.get(naam)
                if b:
                    try:
                        doc.add_picture(io.BytesIO(b), width=Cm(8))
                        doc.add_paragraph(naam).runs[0].font.size = Pt(8)
                    except Exception:  # noqa: BLE001
                        doc.add_paragraph(f"(foto {naam} niet ingevoegd)")
    if v.get("acties"):
        doc.add_heading("Acties", level=1)
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        for c, k in zip(t.rows[0].cells, ("Wie", "Wat", "Tegen wanneer")):
            c.text = k
        for x in v["acties"]:
            cells = t.add_row().cells
            for c, k in zip(cells, ("wie", "wat", "wanneer")):
                c.text = str(x.get(k, ""))
    if v.get("open_punten"):
        doc.add_heading("Open punten (door Mehdi in te vullen)", level=1)
        for x in v["open_punten"]:
            doc.add_paragraph(x, style="List Bullet")
    if s.get("slot"):
        doc.add_paragraph(s["slot"])
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue()


def _tokens(adres):
    return {w for w in re.split(r"[^a-z0-9]+", (adres or "").lower()) if len(w) > 2 and w not in ("belgium", "belgie", "belgië")}


def _tijd(waarde):
    try:
        d = datetime.fromisoformat(str(waarde).replace("Z", "+00:00"))
        return d if d.tzinfo else d.replace(tzinfo=timezone.utc)
    except (ValueError, TypeError):
        return None


def kies_dossiermap(rij, mappen, s):
    """Kies uitsluitend een barstenmap op nummer en bevestigd straat/huisnummer."""
    kandidaten = [(n, p) for n, p in mappen if not re.search(s["overslaan"], n)]
    barsten = [(n, p) for n, p in kandidaten if vs.mapnummer(n) and "barsten" in vs.plat(n)]
    sta = [(n, p) for n, p in kandidaten if re.match(r"^\d{4,6}_STA(?=[_ ]|$)", n)]
    nummer = str(rij.get("dossier") or "").strip()
    hoe_postcode = ""
    if not re.fullmatch(r"\d{4,6}", nummer):
        nummer = ""
    sn = vs.straat_nummer(rij.get("adres"))
    bord_bus = vs.busnummer(rij.get("adres"))
    def postcode(adres):
        staart = str(adres or "").partition(",")[2]
        gevonden = re.findall(r"\b\d{4}\b", staart)
        return gevonden[-1] if gevonden else ""
    def gemeente(adres):
        for deel in str(adres or "").split(",")[1:]:
            deel = re.sub(r"(?i)\b(?:bus|bte|bt)\s*\d+[a-z]?\b", "", deel)
            deel = re.sub(r"(?i)\b(?:b[e]?-)?\d{4}\b", "", deel)
            naam = vs.plat(deel)
            if naam:
                return naam
        return ""
    bord_postcode = postcode(rij.get("adres"))
    bord_gemeente = gemeente(rij.get("adres")) if not bord_postcode else ""
    if len(nummer) == 4 and nummer == bord_postcode:
        nummer, hoe_postcode = "", "nummer is de postcode"
    def straat_in(naam):
        return bool(sn and vs.straat_nummer(str(naam).rsplit("_", 1)[-1]) == sn)
    def plaats_in(naam):
        adresdeel = str(naam).rsplit("_", 1)[-1]
        return bool(straat_in(naam) and
                    (not bord_postcode or postcode(adresdeel) == bord_postcode) and
                    (not bord_gemeente or gemeente(adresdeel) == bord_gemeente))
    def adres_in(naam):
        if not plaats_in(naam):
            return False
        map_bus = vs.busnummer(str(naam).rsplit("_", 1)[-1])
        if bord_bus and map_bus:
            return bord_bus == map_bus
        if bord_bus or map_bus:
            return bool(nummer and vs.mapnummer(naam) == nummer)
        return True
    op_nummer = [(n, p) for n, p in barsten if nummer and vs.mapnummer(n) == nummer]
    op_straat = [(n, p) for n, p in barsten if straat_in(n)]
    op_plaats = [(n, p) for n, p in barsten if plaats_in(n)]
    op_adres = [(n, p) for n, p in barsten if adres_in(n)]
    if len(op_nummer) > 1:
        return None, hoe_postcode, f"meerdere barstenmappen met nummer {nummer}: {', '.join(n for n, _ in op_nummer)}"
    if len(op_adres) > 1:
        if len(op_nummer) == 1 and op_nummer[0] in op_adres:
            op_adres = op_nummer
        else:
            return None, hoe_postcode, f"meerdere barstenmappen op {rij.get('adres')}: {', '.join(n for n, _ in op_adres)}"
    a, b = (op_nummer[0] if op_nummer else None), (op_adres[0] if op_adres else None)
    if a and b and a[1] != b[1]:
        return None, hoe_postcode, (f"agendanummer {nummer} wijst naar {a[0]}, het adres naar {b[0]}; "
                                      "zet de juiste map op de pagina Commandocentrum")
    if b:
        if len(nummer) in (5, 6) and not a:
            return None, hoe_postcode, (f"agendanummer {nummer} staat niet vooraan een barstenmap, maar het adres "
                                          f"raakt {b[0]}; zet de juiste map op de pagina Commandocentrum")
        hoe = "straat en huisnummer"
        if a:
            hoe += "; nummer bevestigt"
        elif nummer:
            hoe += f"; agendanummer {nummer} staat niet vooraan een map (map heeft {vs.mapnummer(b[0])})"
        if hoe_postcode:
            hoe += "; " + hoe_postcode
        return b[1], hoe, None
    if op_plaats:
        return None, hoe_postcode, (f"busnummer van {rij.get('adres')} past niet bij barstenmap(pen) "
                                      f"{', '.join(n for n, _ in op_plaats)}; zet de juiste map op de pagina Commandocentrum")
    if bord_postcode and op_straat:
        return None, hoe_postcode, (f"postcode {bord_postcode} past niet bij barstenmap(pen) "
                                      f"{', '.join(n for n, _ in op_straat)}; zet de juiste map op de pagina Commandocentrum")
    if bord_gemeente and op_straat:
        return None, hoe_postcode, (f"gemeente {bord_gemeente} past niet bij barstenmap(pen) "
                                      f"{', '.join(n for n, _ in op_straat)}; zet de juiste map op de pagina Commandocentrum")
    if a:
        return None, hoe_postcode, (f"agendanummer {nummer} wijst naar {a[0]}, maar het adres bevestigt dat niet; "
                                       "zet de juiste map")
    alleen_sta = [(n, p) for n, p in sta if (nummer and vs.mapnummer(n) == nummer) or adres_in(n)]
    if alleen_sta:
        return None, hoe_postcode, (f"op dit nummer/adres staat alleen de stabiliteitsmap {alleen_sta[0][0]}, zonder de "
                                      "barsten-opbouw; maak de map volgens '1 Mapstructuur_Barsten & Scheuren' "
                                      "of zet de juiste map")
    bezoekjaar = int(str(rij.get("datum") or date.today().isoformat())[:4])
    jaren = sorted({bezoekjaar, bezoekjaar - 1, date.today().year})
    return None, hoe_postcode, (f"geen barstenmap gevonden in {', '.join(str(j) for j in jaren)} "
                                  f"(nummer {nummer}, adres {sn})")


def deel_in(entries, dossiermap, datum, s, contract):
    uit = {k: [] for k in ("fotos", "fotos_iphone", "transcripten", "opnames", "documenten", "mails", "proef", "overig")}
    uit["submappen"] = {k: None for k in s["submappen"]}
    uit["laatst_gewijzigd"] = None
    basis = dossiermap.lower() + "/"
    voorbereid = []
    for e in entries:
        pad = e.get("path_lower") or (e.get("path_display") or "").lower()
        if not pad.startswith(basis):
            continue
        delen = pad[len(basis):].split("/")
        if not delen or any(not d or d.startswith(".") for d in delen):
            continue
        if e.get(".tag") == "folder" and len(delen) == 1:
            sleutel = vs.submap_sleutel(e.get("name") or delen[0])
            if sleutel:
                uit["submappen"][sleutel] = e.get("name") or delen[0]
        elif e.get(".tag") == "file":
            voorbereid.append((e, delen))
    jpg_stammen = {(tuple(d[:-1]), os.path.splitext(d[-1])[0]) for e, d in voorbereid
                   if os.path.splitext(d[-1])[1] in (".jpg", ".jpeg")}
    for e, delen in voorbereid:
        naam = delen[-1]
        ext = os.path.splitext(naam)[1]
        sleutel = vs.submap_sleutel(delen[0])
        if "(concept)" in naam:
            rol = "proef"
        elif sleutel == "fotos" and len(delen) == 3 and (delen[1] == datum or delen[1].startswith(datum + " ")):
            if ext in OPNAME_EXT:
                rol = "opnames"
            elif ext in contract["foto_ext"] and not (ext in (".heic", ".heif") and
                                                              (tuple(delen[:-1]), os.path.splitext(naam)[0]) in jpg_stammen):
                rol = "fotos"
                if delen[1] == datum.lower() + " iphone":
                    uit["fotos_iphone"].append(e)
            else:
                rol = "overig"
        elif sleutel == "plaud" and len(delen) == 2 and pz.is_bezoektranscript(naam, datum):
            rol = "transcripten"
        elif sleutel == "mail" and ext in contract["bijlage_ext"]:
            rol = "mails"
        elif sleutel == "documenten" and ext in contract["bijlage_ext"]:
            rol = "documenten"
        else:
            rol = "overig"
        uit[rol].append(e)
        if e.get("server_modified"):
            lm = e["server_modified"]
            if not uit["laatst_gewijzigd"] or lm > uit["laatst_gewijzigd"]:
                uit["laatst_gewijzigd"] = lm
    return uit


def _vinger(meta, bronnen):
    kern = {"meta": {k: meta[k] for k in ("projectnr", "adres", "bezoekdatum")},
            "bronnen": sorted([[rol, e.get("path_lower") or e.get("path_display", "").lower(),
                                 e.get("content_hash") or e.get("rev") or ""] for rol, e in bronnen])}
    return hashlib.sha256(json.dumps(kern, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _veilig(stam):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", stam).strip("._-")[:80] or "bestand"


def pakket_plan(rij, inv, s, contract):
    dossiermap = rij["dossiermap"]
    mapnr = vs.mapnummer(os.path.basename(dossiermap))
    if not mapnr:
        raise ValueError("dossiermap zonder mapnummer")
    datum = rij["datum"]
    naam = f"{mapnr}_{datum}"
    meta = {"dossier": naam, "projectnr": mapnr, "adres": rij.get("adres") or "", "building": rij.get("adres") or "",
            "bezoekdatum": datum, "klant": rij.get("klant") or "", "bron": {"dossiermap": dossiermap,
            "verslagopdracht": rij.get("id"), "agendanummer": rij.get("dossier") or "", "agenda_titel": rij.get("titel") or ""}}
    overgeslagen = []
    grens = contract["max_bestand_mb"] * 1_000_000
    def kies(rol, entries):
        gekozen = []
        for e in entries:
            if int(e.get("size") or 0) > grens:
                overgeslagen.append({"dropbox": e.get("path_display"), "reden": "bestand groter dan max_bestand_mb"})
            else:
                gekozen.append(e)
        return gekozen
    transcripten = kies("transcript", sorted(inv["transcripten"], key=lambda e: (e.get("name") or "").lower()))
    fotos = kies("foto", sorted(inv["fotos"], key=lambda e: (e.get("path_lower") or "").lower()))
    bijlagen = []
    for rol, groep, voorvoegsel in (("mail", "mails", "mail"), ("document", "documenten", "doc")):
        for i, e in enumerate(kies(rol, sorted(inv[groep], key=lambda x: x.get("path_lower") or "")), 1):
            if i > 99:
                overgeslagen.append({"dropbox": e.get("path_display"), "reden": "meer dan 99 bijlagen van dit type"})
                continue
            stam, ext = os.path.splitext(e.get("name") or "bestand")
            bijlagen.append({"rol": rol, "entry": e, "doel": f"{voorvoegsel}_{i:02d}_{_veilig(stam)}{ext.lower()}"})
    for e in inv.get("overig", []):
        if vs.submap_sleutel((e.get("path_display") or "").replace(dossiermap + "/", "").split("/")[0]) in ("mail", "documenten"):
            overgeslagen.append({"dropbox": e.get("path_display"), "reden": "ongeschikt bestandsformaat"})
    invoerhash = _vinger(meta, [("transcript", e) for e in transcripten] + [("foto", e) for e in fotos])
    bijlagenhash = _vinger(meta, [(b["rol"], b["entry"]) for b in bijlagen])
    totaal_bytes = sum(int(e.get("size") or 0) for e in transcripten + fotos + [b["entry"] for b in bijlagen]) + 1_000_000
    return {"naam": naam, "meta": meta, "transcripten": transcripten, "fotos": fotos, "bijlagen": bijlagen,
            "overgeslagen": overgeslagen, "invoerhash": invoerhash, "bijlagenhash": bijlagenhash, "totaal_bytes": totaal_bytes}


def bruikbaar(rij, inv, nu):
    if not inv["transcripten"] or not inv["fotos"]:
        return False
    einde = _tijd(rij.get("einde"))
    if not inv.get("fotos_iphone") and (not einde or now_aware(nu) <= einde + timedelta(hours=ZONDER_TAKEN_UUR)):
        return False
    if einde:
        return now_aware(nu) > einde
    return str(rij.get("datum") or "") < now_aware(nu).date().isoformat()


def now_aware(nu):
    return nu if nu.tzinfo else nu.replace(tzinfo=timezone.utc)


def dossierkeuze_vast(rij):
    return bool(rij.get("taken") or (rij.get("pakket") or {}).get("overdracht") or
                rij.get("stand") not in (None, "", "gepland", "voorbereid"))


def taken_dicht(rij, nu, lees):
    taken = rij.get("taken") or []
    einde = _tijd(rij.get("einde"))
    oud = bool(einde and now_aware(nu) > einde + timedelta(hours=ZONDER_TAKEN_UUR))
    def open_taak(reden):
        if oud:
            return True, f"bezoek meer dan 6 uur voorbij; {reden}"
        return False, reden

    # Een fototaak die nog niet op het bord staat, bewijst geen volledige levering.
    # Bewaar de laatste versie per wacht: een oude, opgepakte :1 is na :2 historie.
    per_wacht = {}
    zonder_wacht = []
    for taak in taken:
        wacht = taak.get("voor")
        if wacht:
            per_wacht[wacht] = taak
        else:
            zonder_wacht.append(taak)
    if "icloud-wacht" not in per_wacht:
        return open_taak("iCloud-wachttaak nog niet klaargezet")
    gezien = {}
    for taak in zonder_wacht + list(per_wacht.values()):
        try:
            uniek = taak["uniek"]
            if uniek not in gezien:
                try:
                    gezien[uniek] = (True, lees(uniek))
                except Exception as e:  # noqa: BLE001
                    gezien[uniek] = (False, e)
            gelukt, antwoord = gezien[uniek]
            if not gelukt:
                raise antwoord
            items = antwoord.get("items", []) if isinstance(antwoord, dict) else antwoord
            item = next((it for it in items or [] if it.get("uniek") == uniek), None)
            if not item:
                return open_taak("wachttaak ontbreekt op het bord")
            if item.get("status") != "opgepakt":
                return open_taak("iCloud-wachttaak nog open" if taak.get("voor") == "icloud-wacht"
                                 else "wachttaak nog open")
            if taak.get("voor") == "icloud-wacht" and item.get("opgepakt_door") != "icloud-wacht":
                return open_taak("iCloud-levering niet bevestigd")
        except Exception:
            return open_taak("taakstatus niet leesbaar")
    return True, "wachttaken volledig afgehandeld"


def _open_wachttaken(rij, lees, *, overdracht_ts=None, toegestane_taken=None):
    """Geef de open taken met sleutel en reden, eventueel begrensd tot de overdracht."""
    taken = rij.get("taken") or []
    volgorde = list(dict.fromkeys(taak.get("voor") for taak in taken
                                  if taak.get("voor") in ("icloud-wacht", "plaud-wacht")))
    per_wacht = {}
    for taak in reversed(taken):
        wacht = taak.get("voor")
        if wacht not in ("icloud-wacht", "plaud-wacht") or wacht in per_wacht:
            continue
        uniek = taak.get("uniek")
        if toegestane_taken is not None and uniek not in toegestane_taken:
            continue
        if not uniek:
            per_wacht[wacht] = (uniek, "wachttaak zonder sleutel")
            continue
        try:
            antwoord = lees(uniek)
            items = antwoord.get("items", []) if isinstance(antwoord, dict) else antwoord
            item = next((it for it in items or [] if it.get("uniek") == uniek), None)
        except Exception:  # noqa: BLE001
            if overdracht_ts is None:
                per_wacht[wacht] = (uniek, "taakstatus niet leesbaar")
            continue
        if overdracht_ts is not None:
            taak_ts = _tijd((item or {}).get("ts"))
            if not taak_ts or taak_ts > overdracht_ts:
                continue
        if not item or item.get("status") != "opgepakt":
            reden = "iCloud-wachttaak nog open" if wacht == "icloud-wacht" else "Plaud-wachttaak nog open"
        elif wacht == "icloud-wacht" and item.get("opgepakt_door") != "icloud-wacht":
            reden = "iCloud-levering niet bevestigd"
        else:
            reden = None
        per_wacht[wacht] = (uniek, reden)
    return [per_wacht[wacht] for wacht in volgorde
            if wacht in per_wacht and per_wacht[wacht][1]]


def open_wachttaak(rij, lees, *, overdracht_ts=None, toegestane_taken=None):
    """Geef alleen een reden voor een wachttaak die bij deze overdracht meetelt."""
    open_taken = _open_wachttaken(rij, lees, overdracht_ts=overdracht_ts,
                                 toegestane_taken=toegestane_taken)
    return open_taken[0][1] if open_taken else None


def _transcript_tekst(data, naam):
    tekst = wp.docx_tekst(data) if naam.lower().endswith(".docx") else data.decode("utf-8", "replace")
    regels = tekst.splitlines()
    grens = next((i for i, regel in enumerate(regels[:25]) if re.match(r"^\s*[-=_]{20,}\s*$", regel)), None)
    if grens is not None:
        regels = regels[grens + 1:]
    else:
        regels = [regel for i, regel in enumerate(regels) if i >= 15 or not re.match(
            r"(?i)^\s*(plaud[_ -]?id|transcript[_ -]?hash)\s*:", regel)]
    return "\n".join(regels).strip()


def _opnametijd(kop):
    try:
        from PIL import Image
        with Image.open(io.BytesIO(kop)) as img:
            exif = img.getexif()
            raw = (exif.get_ifd(0x8769) or {}).get(36867) or exif.get(306)
        return datetime.strptime(str(raw), "%Y:%m:%d %H:%M:%S").isoformat() if raw else None
    except Exception:
        return None


def toets_pakket(pad, contract):
    """Het gedeelde bestandscontract op een reeds gebouwde tijdelijke map."""
    import stat
    toegestaan = set(contract["wortel"])
    patroon = re.compile(contract["naam_bestand"])
    namen = os.listdir(pad)
    if any(n not in toegestaan for n in namen) or not {"Transcript.docx", "meta.json", "manifest.json"} <= set(namen):
        raise ValueError("onbekend of ontbrekend bestand in pakketwortel")
    totaal, fotos = 0, 0
    for map_pad, soort in ((pad, "wortel"), (os.path.join(pad, "photos"), "foto"),
                           (os.path.join(pad, "bijlagen"), "bijlage")):
        if soort != "wortel" and not os.path.isdir(map_pad):
            continue
        for naam in os.listdir(map_pad):
            if not patroon.fullmatch(naam):
                raise ValueError(f"ongeldige bestandsnaam: {naam}")
            volledig = os.path.join(map_pad, naam)
            mode = os.lstat(volledig).st_mode
            if stat.S_ISLNK(mode):
                raise ValueError("symlink in pakket")
            if stat.S_ISDIR(mode):
                if soort != "wortel" or naam not in ("photos", "bijlagen"):
                    raise ValueError("onverwachte submap in pakket")
                continue
            if not stat.S_ISREG(mode):
                raise ValueError("onbekend bestandstype in pakket")
            ext = os.path.splitext(naam)[1].lower()
            if soort == "foto":
                if ext not in contract["foto_ext"]:
                    raise ValueError("ongeldig fotoformaat")
                fotos += 1
            if soort == "bijlage" and (ext not in contract["bijlage_ext"] or
                                        not re.match(r"^(mail|doc)_\d{2}_", naam)):
                raise ValueError("ongeldig bijlageformaat")
            grootte = os.stat(volledig).st_size
            if grootte > contract["max_bestand_mb"] * 1_000_000:
                raise ValueError("bestand boven pakketgrens")
            if naam == "meta.json" and soort == "wortel" and grootte > contract["meta_max_bytes"]:
                raise ValueError("meta.json te groot")
            if naam == "manifest.json" and soort == "wortel" and grootte > contract["manifest_max_bytes"]:
                raise ValueError("manifest.json te groot")
            totaal += grootte
    if not 1 <= fotos <= contract["max_fotos"]:
        raise ValueError("fotoaantal buiten grens")
    if totaal > contract["max_totaal_mb"] * 1_000_000:
        raise ValueError("pakket boven totaalgrens")
    with open(os.path.join(pad, "meta.json"), encoding="utf-8") as f:
        meta = json.load(f)
    if set(meta) != set(contract["meta_sleutels"]) or set(meta["bron"]) != set(contract["bron_sleutels"]):
        raise ValueError("meta.json wijkt af van contract")
    return True


def schrijf_pakket(plan, inbox, download, download_kop, nu):
    """Schrijf één pakket atomair; alle opschoning blijft binnen eigen tijdelijke mappen."""
    contract = vs.pakketcontract()
    os.makedirs(inbox, mode=0o755, exist_ok=True)
    st = os.stat(inbox)
    if st.st_uid != os.getuid() or not os.access(inbox, os.W_OK):
        raise PermissionError(f"inbox {inbox} niet bruikbaar; moet van ubuntu zijn met rechten 755")
    lockpad = os.path.join(inbox, ".lock")
    try:
        fd = os.open(lockpad, os.O_CREAT | os.O_EXCL | os.O_RDWR, 0o644)
    except FileExistsError:
        fd = os.open(lockpad, os.O_RDWR)
    tmp = None
    try:
        fcntl.flock(fd, fcntl.LOCK_EX)
        for naam in os.listdir(inbox):
            pad = os.path.join(inbox, naam)
            if re.match(r"^\.\d{4,6}_\d{4}-\d{2}-\d{2}(-\d+)?\.tmp-", naam) and \
                    os.path.isdir(pad) and not os.path.islink(pad) and \
                    now_aware(nu).timestamp() - os.stat(pad).st_mtime > TMP_OUD_UUR * 3600:
                shutil.rmtree(pad)
                print(f"verweesde tijdelijke pakketmap verwijderd: {naam}", file=sys.stderr)
        versies = []
        for naam in os.listdir(inbox):
            if re.fullmatch(re.escape(plan["naam"]) + r"(?:-\d+)?", naam) and os.path.isdir(os.path.join(inbox, naam)):
                versies.append(naam)
        for naam in versies:
            with open(os.path.join(inbox, naam, "manifest.json"), encoding="utf-8") as f:
                bestaand = json.load(f)
            if bestaand.get("invoerhash") == plan["invoerhash"]:
                return {"uitkomst": "al klaargezet", "naam": naam, "pad": os.path.join(inbox, naam),
                        "invoerhash": plan["invoerhash"], "bijlagenhash": bestaand.get("bijlagenhash"),
                        "ts": bestaand.get("aangemaakt")}
        if len(versies) >= MAX_VERSIES:
            return {"uitkomst": "grens"}
        doelnaam = plan["naam"]
        i = 2
        while os.path.lexists(os.path.join(inbox, doelnaam)):
            doelnaam = f"{plan['naam']}-{i}"
            i += 1
        tmp = tempfile.mkdtemp(prefix=f".{doelnaam}.tmp-", dir=inbox)
        os.chmod(tmp, 0o755)
        os.mkdir(os.path.join(tmp, "photos"), 0o755)
        os.mkdir(os.path.join(tmp, "bijlagen"), 0o755)
        manifest_bronnen = []
        gekozen = {}
        for e in plan["transcripten"]:
            kop = download_kop(e["path_display"], 2000).decode("utf-8", "replace")
            idregel = next((v.strip() for k, _, v in (regel.partition(":") for regel in kop.splitlines()[:10])
                             if k.strip().lower().replace("-", "_").replace(" ", "_") == "plaud_id"), "")
            sleutel = idregel or e["path_display"]
            if sleutel not in gekozen or (e.get("server_modified") or "") > (gekozen[sleutel].get("server_modified") or ""):
                gekozen[sleutel] = e
        from docx import Document
        doc = Document()
        for i, e in enumerate(sorted(gekozen.values(), key=lambda x: (x.get("name") or "").lower()), 1):
            data = download(e["path_display"])
            if e.get("size") is not None and len(data) != e["size"]:
                raise ValueError("transcriptlengte wijkt af van Dropbox-lijst")
            doc.add_paragraph(f"Opname {i}: {e['name']}")
            for regel in _transcript_tekst(data, e["name"]).splitlines():
                doc.add_paragraph(regel)
            manifest_bronnen.append({"rol": "transcript", "dropbox": e["path_display"], "rev": e.get("rev"),
                                     "content_hash": e.get("content_hash"), "bytes": len(data),
                                     "doel": "Transcript.docx", "opnametijd": None})
        doc_data = io.BytesIO()
        doc.save(doc_data)
        with open(os.path.join(tmp, "Transcript.docx"), "xb") as f:
            f.write(doc_data.getvalue())
        foto_tijden = [(e, _opnametijd(download_kop(e["path_display"], 131072))) for e in plan["fotos"]]
        datum = plan["meta"]["bezoekdatum"].casefold()
        def foto_volgorde(foto):
            e, opnametijd = foto
            mapnaam = os.path.basename(os.path.dirname(e["path_display"])).casefold()
            groep = (0, "") if mapnaam == datum + " iphone" else \
                    (1, mapnaam) if mapnaam == datum or mapnaam.startswith(datum + " ") else (2, "")
            return groep, opnametijd is None, opnametijd or "", e.get("path_lower") or e["path_display"].casefold()
        foto_tijden.sort(key=foto_volgorde)
        for i, (e, opnametijd) in enumerate(foto_tijden, 1):
            stam, ext = os.path.splitext(e["name"])
            doel = f"{i:03d}_{_veilig(stam)}{ext.lower()}"
            data = download(e["path_display"])
            if e.get("size") is not None and len(data) != e["size"]:
                raise ValueError("fotolengte wijkt af van Dropbox-lijst")
            with open(os.path.join(tmp, "photos", doel), "xb") as f:
                f.write(data)
            manifest_bronnen.append({"rol": "foto", "dropbox": e["path_display"], "rev": e.get("rev"),
                                     "content_hash": e.get("content_hash"), "bytes": len(data),
                                     "doel": "photos/" + doel, "opnametijd": opnametijd})
        for b in plan["bijlagen"]:
            e = b["entry"]
            data = download(e["path_display"])
            if e.get("size") is not None and len(data) != e["size"]:
                raise ValueError("bijlagelengte wijkt af van Dropbox-lijst")
            with open(os.path.join(tmp, "bijlagen", b["doel"]), "xb") as f:
                f.write(data)
            manifest_bronnen.append({"rol": b["rol"], "dropbox": e["path_display"], "rev": e.get("rev"),
                                     "content_hash": e.get("content_hash"), "bytes": len(data),
                                     "doel": "bijlagen/" + b["doel"], "opnametijd": None})
        with open(os.path.join(tmp, "meta.json"), "x", encoding="utf-8") as f:
            json.dump(plan["meta"], f, ensure_ascii=False, indent=2)
        manifest = {"versie": 1, "contract": 1, "naam": doelnaam, "invoerhash": plan["invoerhash"],
                    "bijlagenhash": plan["bijlagenhash"], "aangemaakt": now_aware(nu).isoformat(),
                    "bronnen": manifest_bronnen, "overgeslagen": plan["overgeslagen"]}
        with open(os.path.join(tmp, "manifest.json"), "x", encoding="utf-8") as f:
            json.dump(manifest, f, ensure_ascii=False, indent=2)
        toets_pakket(tmp, contract)
        for map_pad, _, bestanden in os.walk(tmp):
            for naam in bestanden:
                os.chmod(os.path.join(map_pad, naam), 0o644)
        doel = os.path.join(inbox, doelnaam)
        if os.path.lexists(doel):
            raise FileExistsError(doel)
        os.rename(tmp, doel)
        tmp = None
        return {"uitkomst": "klaargezet", "naam": doelnaam, "pad": doel, "invoerhash": plan["invoerhash"],
                "bijlagenhash": plan["bijlagenhash"], "ts": manifest["aangemaakt"], "fotos": len(foto_tijden),
                "bijlagen": len(plan["bijlagen"]), "transcripten": len(gekozen), "overgeslagen": plan["overgeslagen"]}
    finally:
        if tmp and os.path.isdir(tmp):
            shutil.rmtree(tmp)
        fcntl.flock(fd, fcntl.LOCK_UN)
        os.close(fd)


def _bijlagen_na_overdracht(plan, overdracht):
    try:
        with open(os.path.join(overdracht["pad"], "manifest.json"), encoding="utf-8") as f:
            manifest = json.load(f)
        oud = {(b["rol"], b["dropbox"].lower(), b.get("content_hash") or b.get("rev"))
               for b in manifest.get("bronnen", []) if b.get("rol") in ("mail", "document")}
    except (OSError, KeyError, ValueError):
        oud = set()
    return sum((b["rol"], b["entry"]["path_display"].lower(),
                b["entry"].get("content_hash") or b["entry"].get("rev")) not in oud for b in plan["bijlagen"])


class VerslagAgent:
    def __init__(self, soort):
        self.soort = soort
        self.s = vs.SOORTEN[soort]
        self.naam = self.s["agent"]
        self.ag = bord.Agent(self.naam)
        self._mapcache = {}
        self._opbouwcache = {}
        self._inventaris = {}
        self._taakstatus = {}
        self._zoeknood = None
        self._noden = []

    # ------------------------------------------------------------- bord ---
    def rijen(self, alle=False):
        return bord.call(f"/api/verslagopdracht?agent={self.naam}" + ("&open=alle" if alle else "")).get("rijen", [])

    def rij(self, uniek=None, vid=None):
        q = f"uniek={uniek}" if uniek else f"id={vid}"
        r = bord.call(f"/api/verslagopdracht?open=alle&{q}").get("rijen", [])
        return r[0] if r else None

    def bewaar(self, uniek, **velden):
        velden["_alleen"] = list(velden)
        bord.call("/api/verslagopdracht", {"rijen": [{"uniek": uniek, **velden}]})

    # ------------------------------------------------------- dossiermap ---
    def tkn_mappen(self, rij):
        basis = self.s["basismappen"][0]
        jaar = int(str(rij.get("datum") or date.today().isoformat())[:4])
        uit = []
        for j in sorted({jaar, jaar - 1, date.today().year}):
            pad = f"{basis}/{j}"
            if pad not in self._mapcache:
                self._mapcache[pad] = [(e.get("name") or "", e.get("path_display") or "") for e in
                                       (bronnen.lijst(pad, recursief=False) or []) if e.get(".tag") == "folder"]
            uit.extend(self._mapcache[pad])
        return uit

    def opbouw(self, pad):
        if pad not in self._opbouwcache:
            entries = bronnen.lijst(pad, recursief=False)
            self._opbouwcache[pad] = None if entries is None else {
                vs.submap_sleutel(e.get("name") or ""): e.get("name") for e in entries
                if e.get(".tag") == "folder" and vs.submap_sleutel(e.get("name") or "")}
        return self._opbouwcache[pad]

    def zoek_dossiermap(self, rij):
        """Eerst het nummer vooraan de mapnaam, dan het adres (straat + huisnummer), dan de klantnaam."""
        if self.s.get("jaarmappen"):
            pad, hoe, nood = kies_dossiermap(rij, self.tkn_mappen(rij), self.s)
            self._zoeknood = nood
            if pad:
                opbouw = self.opbouw(pad)
                ontbreekt = [self.s["submappen"][k] for k in self.s["vereiste_submappen"] if not opbouw or k not in opbouw]
                if ontbreekt:
                    self._zoeknood = (f"map {os.path.basename(pad)} mist {', '.join(ontbreekt)}; maak de submappen "
                                      "volgens het sjabloon (agents maken geen submappen)")
                    return None, ""
            return pad, hoe
        nummer = str(rij.get("dossier") or "").strip()
        adres_t = _tokens(rij.get("adres"))
        klant_t = _tokens(rij.get("klant"))
        kandidaten = []
        for basis in self.s["basismappen"]:
            for e in bronnen.lijst(basis, recursief=False) or []:
                if e.get(".tag") != "folder":
                    continue
                naam = e.get("name", "")
                if naam[:1] in "0_" and not re.match(r"^\d{4,5}\s", naam):
                    continue   # 0. Archive, 0. Structure, ...
                kandidaten.append((naam, e.get("path_display")))
        if nummer:
            for naam, pad in kandidaten:
                if re.match(rf"^\s*{re.escape(nummer)}\b", naam):
                    return pad, "projectnummer vooraan de mapnaam"
        beste, score, hoe = None, 0, ""
        for naam, pad in kandidaten:
            nt = _tokens(naam)
            sa = len(nt & adres_t)
            if adres_t and sa >= 2 and sa > score:
                beste, score, hoe = pad, sa, f"adres ({sa} woorden)"
        if beste:
            return beste, hoe
        for naam, pad in kandidaten:
            nt = _tokens(naam)
            sk = len(nt & klant_t)
            if klant_t and sk >= 2 and sk > score:
                beste, score, hoe = pad, sk, f"klantnaam ({sk} woorden)"
        return (beste, hoe) if beste else (None, "")

    def maak_bezoekmap(self, rij, dossiermap):
        if self.s.get("bezoekmap_is_dossiermap"):
            return dossiermap, False
        pad = f"{dossiermap}/{self.s['communicatiemap']}/{rij['datum']} {self.s['bezoeknaam']}"
        if bronnen.lijst(pad, recursief=False) is not None:
            return pad, False
        kop = [f"# {rij['datum']} {self.s['bezoeknaam']}", "",
               f"Bezoekmap aangemaakt door {self.naam} (Mehdi Agents) op {date.today().isoformat()}, op impuls van het Commandocentrum.", "",
               f"- Verslagsoort: {self.s['label']}", f"- Dossier: {rij.get('dossier') or '(geen nummer in de agenda)'}",
               f"- Klant: {rij.get('klant') or ''}", f"- Adres: {rij.get('adres') or ''}",
               f"- Agenda: {rij.get('agenda')} - {rij.get('titel')}", f"- Venster: {rij.get('start')} tot {rij.get('einde')}",
               f"- Deal: {rij.get('deal_id') or '-'}", "",
               "Wat hier landt: `fotos/` (iCloud-wacht, foto's van de bezoekdag binnen 300 m), `transcript.txt` (Plaudwacht),",
               "en daarna het concept-verslag (markdown en Word). Niets wordt overschreven.", ""]
        bronnen.upload(f"{pad}/00 bezoek.md", "\n".join(kop).encode("utf-8"))
        return pad, True

    # ----------------------------------------------------------- pakket ---
    def meet_pakket(self, rij):
        if self.s.get("submappen"):
            pk = {"agenda": True, "dossiermap": bool(rij.get("dossiermap")), "bezoekmap": bool(rij.get("bezoekmap")),
                  "fotos": 0, "fotos_iphone": 0, "opnames": 0, "transcripten": 0, "documenten": 0, "mails": 0,
                  "proef": 0, "submappen": {k: None for k in self.s["submappen"]},
                  "ontbrekende_submappen": list(self.s["vereiste_submappen"]), "laatst_gewijzigd": None,
                  "gemeten_ts": datetime.now(timezone.utc).isoformat()}
            if not rij.get("dossiermap"):
                return pk
            entries = bronnen.lijst(rij["dossiermap"], recursief=True)
            if entries is None:
                pk["dossiermap"] = pk["bezoekmap"] = False
                return pk
            inv = deel_in(entries, rij["dossiermap"], rij["datum"], self.s, vs.pakketcontract())
            self._inventaris[rij["uniek"]] = inv
            pk.update({k: len(inv[k]) for k in ("fotos", "fotos_iphone", "opnames", "transcripten", "documenten", "mails", "proef")})
            pk["submappen"] = inv["submappen"]
            pk["ontbrekende_submappen"] = [k for k in self.s["vereiste_submappen"] if not inv["submappen"].get(k)]
            pk["laatst_gewijzigd"] = inv["laatst_gewijzigd"]
            return pk
        pk = {"agenda": True, "dossiermap": bool(rij.get("dossiermap")), "bezoekmap": bool(rij.get("bezoekmap")),
              "fotos": 0, "opnames": 0, "transcripten": 0, "documenten": 0, "proef": 0, "gemeten_ts": datetime.now().isoformat(timespec="minutes")}
        if not rij.get("bezoekmap"):
            return pk
        items = bronnen.lijst(rij["bezoekmap"], recursief=True)
        if items is None:
            pk["bezoekmap"] = False
            return pk
        for e in items:
            if e.get(".tag") != "file":
                continue
            n = e.get("name", "").lower()
            if n.endswith(FOTO_EXT):
                pk["fotos"] += 1
            elif n.endswith(OPNAME_EXT):
                pk["opnames"] += 1
            elif "transcript" in n:
                pk["transcripten"] += 1
            elif "(concept)" in n or "proef" in n:
                pk["proef"] += 1
            elif n.endswith(DOC_EXT) and n != "00 bezoek.md" and n != "00 fotos.md":
                pk["documenten"] += 1
        return pk

    def _bewaar_pakket(self, rij, pk, **velden):
        oud = rij.get("pakket") or {}
        for k in ("overdracht", "kandidaat", "dossierkeuze", "bijlagen_na_overdracht", "versiegrens"):
            if k not in pk and k in oud:
                pk[k] = oud[k]
        self.bewaar(rij["uniek"], pakket=pk, **velden)
        rij["pakket"] = pk
        return pk

    # --------------------------------------------------------- impulsen ---
    def impulsen(self):
        items = [it for it in bord.klaargezet_voor(self.naam, n=50) if it.get("soort") == "impuls"]
        gedaan = 0
        for it in reversed(items):
            try:
                d = json.loads(it.get("inhoud") or "{}")
            except ValueError:
                d = {}
            uniek, fase = d.get("uniek"), d.get("fase")
            rij = self.rij(uniek=uniek) if uniek else None
            onderwerp = f"{d.get('datum','?')} {d.get('dossier') or d.get('klant') or ''}".strip()
            if not rij:
                self.ag.log(onderwerp, "fout", f"impuls zonder rij op het bord: {it.get('titel','')[:80]}")
                bord.opgepakt(it["id"], self.naam)
                continue
            try:
                if fase == "voorbereiding":
                    self.voorbereiding(rij, onderwerp)
                elif fase == "verslag":
                    pk = self.meet_pakket(rij)
                    self._bewaar_pakket(rij, pk, stand="pakket klaar" if not rij.get("proef_pad") else rij["stand"])
                    self.ag.log(onderwerp, "bevinding", f"pakket vol: {pk['fotos']} foto's, {pk['transcripten']} transcript(en), "
                                f"{pk['documenten']} document(en); klaar voor verslag, wacht op de knop Proef maken")
                gedaan += 1
            except Exception as e:  # noqa: BLE001
                self.ag.log(onderwerp, "fout", f"impuls {fase} mislukt: {type(e).__name__}: {str(e)[:200]}")
            bord.opgepakt(it["id"], self.naam)
        return gedaan

    def voorbereiding(self, rij, onderwerp):
        if self.s.get("submappen"):
            uniek = rij["uniek"]
            pad = rij.get("dossiermap") or ""
            hoe = "door Mehdi gezet"
            if not pad:
                pad, hoe = self.zoek_dossiermap(rij)
            if not pad:
                self._noden.append({"tekst": self._zoeknood or f"{onderwerp}: dossiermap niet gevonden", "wie": "mehdi"})
                self._bewaar_pakket(rij, {"agenda": True, "dossiermap": False, "bezoekmap": False}, stand="voorbereid")
                return
            if not vs.tkn_pad(pad, self.s["basismappen"][0]):
                self._noden.append({"tekst": "dossiermap op het bord heeft niet de vorm /Work/TKN BURO WORK/1 Projects/<jaar>/<nr>_...; zet ze opnieuw op de pagina Commandocentrum", "wie": "mehdi"})
                return
            opbouw = self.opbouw(pad)
            ontbreekt = [self.s["submappen"][k] for k in self.s["vereiste_submappen"] if not opbouw or k not in opbouw]
            if ontbreekt:
                self._nood(f"map {os.path.basename(pad)} mist {', '.join(ontbreekt)}; maak de submappen volgens het sjabloon (agents maken geen submappen)")
                return
            bezoekmap, _ = self.maak_bezoekmap(rij, pad)
            pk = self.meet_pakket({**rij, "dossiermap": pad, "bezoekmap": bezoekmap})
            if not rij.get("dossiermap"):
                pk["dossierkeuze"] = {"door": "agent", "pad": pad, "hoe": hoe, "ts": datetime.now(timezone.utc).isoformat()}
            self._bewaar_pakket(rij, pk, dossiermap=pad, bezoekmap=bezoekmap, stand="voorbereid")
            self.ag.log(onderwerp, "schrijf", f"dossiermap gevonden ({hoe}); de dossiermap is de bezoekmap")
            return
        uniek = rij["uniek"]
        dossiermap = rij.get("dossiermap") or ""
        hoe = "door Mehdi gezet"
        if not dossiermap:
            dossiermap, hoe = self.zoek_dossiermap(rij)
        if not dossiermap:
            self.ag.log(onderwerp, "bevinding", f"geen dossiermap gevonden in {len(self.s['basismappen'])} basismap(pen) "
                        f"(nummer '{rij.get('dossier')}', adres '{rij.get('adres')}', klant '{rij.get('klant')}')")
            self.bewaar(uniek, pakket={"agenda": True, "dossiermap": False, "bezoekmap": False}, stand="voorbereid")
            self._noden.append({"tekst": f"{onderwerp}: dossiermap niet gevonden; zet het pad op de pagina Commandocentrum", "wie": "mehdi"})
            return
        bezoekmap, nieuw = self.maak_bezoekmap(rij, dossiermap)
        rij = {**rij, "dossiermap": dossiermap, "bezoekmap": bezoekmap}
        pk = self.meet_pakket(rij)
        self._bewaar_pakket(rij, pk, dossiermap=dossiermap, bezoekmap=bezoekmap, stand="voorbereid")
        self.ag.log(onderwerp, "schrijf", f"dossiermap gevonden ({hoe}); bezoekmap {'aangemaakt' if nieuw else 'bestond al'}: {bezoekmap}")

    # -------------------------------------------------------- opdrachten ---
    def opdrachten(self):
        items = [it for it in bord.klaargezet_voor(self.naam, n=30) if it.get("soort") == "opdracht"]
        gedaan = 0
        for it in reversed(items):
            m = re.match(r"proef\s+(\d+)", it.get("titel", ""))
            # Regel van Mehdi (13-09-2026): alleen opdrachten van de knop op het bord kosten tokens.
            if not m or not str(it.get("van", "")).startswith("bord:"):
                self.ag.log(it.get("sleutel", "?"), "besluit", f"opdracht overgeslagen (niet van de knop): {it.get('titel','')} van {it.get('van','')}")
                bord.opgepakt(it["id"], self.naam)
                continue
            rij = self.rij(vid=int(m.group(1)))
            onderwerp = f"{rij.get('datum')} {rij.get('dossier') or rij.get('klant') or ''}".strip() if rij else m.group(1)
            try:
                if not rij:
                    raise RuntimeError("rij niet gevonden op het bord")
                self.proef(rij, onderwerp)
                gedaan += 1
            except Exception as e:  # noqa: BLE001
                self.ag.log(onderwerp, "fout", f"proef mislukt: {type(e).__name__}: {str(e)[:300]}")
                if rij:
                    self.bewaar(rij["uniek"], stand="pakket klaar")
            bord.opgepakt(it["id"], self.naam)
        return gedaan

    def lees_dossiermap(self, rij, inv):
        teksten = []
        for rol, sleutel in (("transcript", "transcripten"), ("mail", "mails"), ("document", "documenten")):
            for e in sorted(inv[sleutel], key=lambda x: x.get("path_lower") or ""):
                naam, pad = e["name"], e["path_display"]
                ext = os.path.splitext(naam.lower())[1]
                if ext in (".msg", ".xlsx", ".pptx"):
                    teksten.append({"naam": naam, "pad": pad, "tekst": "", "leesfout": f"{ext} overgeslagen", "rol": rol})
                    continue
                try:
                    data = bronnen.download(pad, maximum=MAX_BESTAND)
                    if ext == ".pdf":
                        tekst = wp.pdf_tekst(data)
                        if not tekst.strip():
                            raise ValueError("PDF zonder tekst")
                    elif ext == ".docx":
                        tekst = wp.docx_tekst(data)
                    else:
                        tekst = data.decode("utf-8", "replace")
                    teksten.append({"naam": naam, "pad": pad, "tekst": tekst, "rol": rol})
                except Exception as e2:
                    teksten.append({"naam": naam, "pad": pad, "tekst": "", "leesfout": type(e2).__name__, "rol": rol})
        fotos = [{"naam": e["name"], "pad": e["path_display"]} for e in inv["fotos"]]
        return teksten, fotos

    def proef(self, rij, onderwerp):
        if not rij.get("bezoekmap"):
            raise RuntimeError("geen bezoekmap; eerst de voorbereiding (dossiermap)")
        if self.s.get("submappen"):
            inv = self._inventaris.get(rij["uniek"])
            if inv is None:
                entries = bronnen.lijst(rij["dossiermap"], recursief=True)
                if entries is None:
                    raise RuntimeError("dossiermap niet gevonden")
                inv = deel_in(entries, rij["dossiermap"], rij["datum"], self.s, vs.pakketcontract())
            documentenmap = inv["submappen"].get("documenten")
            if not documentenmap:
                raise RuntimeError("documentenmap ontbreekt; maak ze volgens het sjabloon")
            teksten, fotos = self.lees_dossiermap(rij, inv)
            opnames = inv["opnames"]
            rij = {**rij, "_projectnummer": vs.mapnummer(os.path.basename(rij["dossiermap"]))}
            maximum = self.s["bronnen_max_tekens"]
            bron = list(teksten)
            while True:
                try:
                    bundel_tekst = wp.bronnenbundel(bron, maximum=maximum)
                    break
                except wp.extractie.BronLeesFout:
                    weg = next((i for i in range(len(bron) - 1, -1, -1) if bron[i]["rol"] == "document"), None)
                    if weg is None:
                        weg = next((i for i in range(len(bron) - 1, -1, -1) if bron[i]["rol"] == "mail"), None)
                    if weg is None:
                        raise RuntimeError("transcript alleen al boven de bronlimiet")
                    self.ag.log(onderwerp, "bevinding", f"bron buiten proefbundel: {bron[weg]['naam']}")
                    bron.pop(weg)
            teksten = bron
        else:
            teksten, fotos, opnames = wp.lees_bezoekmap(rij["bezoekmap"])
            bundel_tekst = None
        self.ag.log(onderwerp, "bron", f"bezoekmap gelezen: {len(teksten)} tekstbron(nen), {len(fotos)} foto's, {len(opnames)} opname(s)",
                    [t["naam"] + f" ({len(t['tekst'])} tekens)" for t in teksten])
        if bundel_tekst is None:
            bundel_tekst = wp.bronnenbundel(teksten)
        kop = (f"{self.s['label']} - dossier {rij.get('_projectnummer') or rij.get('dossier') or '(geen nummer)'}, {rij.get('adres','')}, klant {rij.get('klant','')}. "
               f"Bezoek op {rij['datum']} van {(rij.get('start') or '')[11:16]} tot {(rij.get('einde') or '')[11:16]}. "
               f"Bezoekmap: {rij['bezoekmap']}." + (f"\nOpmerking van Mehdi: {rij['opmerking']}" if rij.get("opmerking") else ""))
        if rij.get("_projectnummer") and str(rij.get("dossier") or "") != rij["_projectnummer"]:
            kop += f"\nAgendanummer: {rij.get('dossier') or '-'}"
        bundel = bundel_tekst + wp.fotolijst(fotos)
        resp = wp._client().messages.create(model=wp.MODEL, max_tokens=16000, system=_systeem(self.s),
                                            messages=[{"role": "user", "content": kop + "\n\n## BRONNEN\n" + bundel}],
                                            tools=[VERSLAG_SCHEMA], tool_choice={"type": "tool", "name": "verslag"})
        v = wp._tool(resp)
        afgekapt = resp.stop_reason == "max_tokens"
        beelden = {}
        gevraagd = {n for h in v.get("hoofdstukken") or [] for p in h.get("punten") or [] for n in (p.get("fotos") or [])}
        for f in fotos:
            if f["naam"] in gevraagd and f["naam"].lower().endswith((".jpg", ".jpeg", ".png")) and len(beelden) < 24:
                try:
                    beelden[f["naam"]] = bronnen.download(f["pad"])
                except Exception:  # noqa: BLE001
                    pass
        md = _md(v, self.s, rij)
        docx = _docx(v, self.s, rij, beelden)
        naam = f"{self.s.get('verslagtitel', self.s['label'])} {rij.get('_projectnummer') or rij.get('dossier') or rij.get('klant') or ''} {rij['datum']} (concept)".strip()
        doelmap = f"{rij['dossiermap']}/{documentenmap}" if self.s.get("submappen") else rij["bezoekmap"]
        pad_docx = bronnen.upload(f"{doelmap}/{naam}.docx", docx)
        bronnen.upload(f"{doelmap}/{naam}.md", md.encode("utf-8"))
        n_punten = sum(len(h.get("punten") or []) for h in v.get("hoofdstukken") or [])
        info = {"hoofdstukken": len(v.get("hoofdstukken") or []), "punten": n_punten, "fotos": len(beelden),
                "open": len(v.get("open_punten") or []), "docx_kb": len(docx) // 1024, "model": wp.MODEL,
                "tokens": f"{resp.usage.input_tokens}+{resp.usage.output_tokens}", "afgekapt": afgekapt}
        self.bewaar(rij["uniek"], proef_pad=pad_docx, proef_ts=date.today().isoformat(), verslag_md=md, proef_info=info, stand="proef klaar")
        self.ag.log(onderwerp, "proef" if not afgekapt else "fout",
                    f"proef gezet: {os.path.basename(pad_docx)} ({info['docx_kb']} kB, {n_punten} punten, {len(beelden)} foto's, "
                    f"{info['open']} open punten; {info['tokens']} tokens{'; AFGEKAPT op max_tokens' if afgekapt else ''})")

    def _nood(self, tekst):
        self._noden.append({"tekst": tekst, "wie": "mehdi"})

    def _lees_taakstatus(self, uniek):
        if uniek not in self._taakstatus:
            try:
                self._taakstatus[uniek] = (True, bord.call("/api/klaarzet?" + urlencode(
                    {"uniek": uniek, "status": "alle", "n": 5})))
            except Exception as e:  # noqa: BLE001
                self._taakstatus[uniek] = (False, e)
        ok, waarde = self._taakstatus[uniek]
        if not ok:
            raise waarde
        return waarde

    def _ronde_dossiermap(self, rij):
        uniek = rij["uniek"]
        onderwerp = f"{rij.get('datum')} {rij.get('dossier') or rij.get('klant') or ''}".strip()
        basis = self.s["basismappen"][0]
        pad = rij.get("dossiermap") or ""
        oud = rij.get("pakket") or {}
        oude_waarschuwing = (oud.get("overdracht") or {}).get("waarschuwing")
        waarschuwing_nood = {"tekst": oude_waarschuwing, "wie": "mehdi"} if oude_waarschuwing else None
        if waarschuwing_nood:
            self._noden.insert(0, waarschuwing_nood)
        keuze = oud.get("dossierkeuze") or {}
        agentkeuze = keuze.get("door") == "agent" and keuze.get("pad") == pad
        if not pad or (agentkeuze and not dossierkeuze_vast(rij)):
            nieuw, hoe = self.zoek_dossiermap(rij)
            if not nieuw:
                if pad and agentkeuze:
                    self.bewaar(uniek, dossiermap="", bezoekmap="")
                self._nood(self._zoeknood or f"{onderwerp}: geen barstenmap gevonden")
                return
            if nieuw != pad:
                dossierkeuze = {"door": "agent", "pad": nieuw, "hoe": hoe, "ts": datetime.now(timezone.utc).isoformat()}
                oud = {**oud, "dossierkeuze": dossierkeuze}
                rij["pakket"] = oud
                self._bewaar_pakket(rij, oud, dossiermap=nieuw, bezoekmap=nieuw,
                                    **({"stand": "voorbereid"} if not rij.get("stand") or rij.get("stand") == "gepland" else {}))
                self.ag.log(onderwerp, "schrijf", f"dossiermap herzien: {pad} -> {nieuw} ({hoe})" if pad else
                            f"dossiermap gevonden ({hoe}); de dossiermap is de bezoekmap")
                pad = nieuw
        if not vs.tkn_pad(pad, basis):
            self._nood("dossiermap op het bord heeft niet de vorm /Work/TKN BURO WORK/1 Projects/<jaar>/<nr>_...; zet ze opnieuw op de pagina Commandocentrum")
            return
        opbouw = self.opbouw(pad)
        ontbreekt = [self.s["submappen"][k] for k in self.s["vereiste_submappen"] if not opbouw or k not in opbouw]
        if ontbreekt:
            self._nood(f"map {os.path.basename(pad)} mist {', '.join(ontbreekt)}; maak de submappen volgens het sjabloon (agents maken geen submappen)")
            return
        if rij.get("bezoekmap") != pad:
            self.bewaar(uniek, bezoekmap=pad)
        rij = {**rij, "dossiermap": pad, "bezoekmap": pad}
        pk = self.meet_pakket(rij)
        inv = self._inventaris.get(uniek)
        if not inv:
            self._nood(f"{onderwerp}: dossiermap niet leesbaar")
            return
        for sleutel in pk["ontbrekende_submappen"]:
            self._nood(f"map {os.path.basename(pad)} mist {self.s['submappen'][sleutel]}; maak de submappen volgens het sjabloon")
        nu = datetime.now(timezone.utc)
        mapnr = vs.mapnummer(os.path.basename(pad))
        einde = _tijd(rij.get("einde"))
        if einde and nu > einde + timedelta(hours=2):
            if not inv["transcripten"]:
                self._nood(f"{rij['datum']} {mapnr}: nog geen transcript in {inv['submappen'].get('plaud') or self.s['submappen']['plaud']}")
            if not inv["fotos"]:
                self._nood(f"{rij['datum']} {mapnr}: nog geen foto's in {inv['submappen'].get('fotos') or self.s['submappen']['fotos']}/{rij['datum']} ...")
        try:
            plan = pakket_plan(rij, inv, self.s, vs.pakketcontract())
            overdracht = oud.get("overdracht") or {}
            if overdracht.get("invoerhash") == plan["invoerhash"]:
                if overdracht.get("bijlagenhash") != plan["bijlagenhash"]:
                    nieuw_aantal = _bijlagen_na_overdracht(plan, overdracht)
                    if pk.get("bijlagen_na_overdracht") != nieuw_aantal and oud.get("bijlagen_na_overdracht") != nieuw_aantal:
                        self.ag.log(onderwerp, "bevinding", "nieuwe mail of document na de overdracht; staat in Dropbox, niet in het pakket")
                    pk["bijlagen_na_overdracht"] = nieuw_aantal
            else:
                if (oud.get("kandidaat") or {}).get("hash") != plan["invoerhash"]:
                    pk["kandidaat"] = {"hash": plan["invoerhash"], "ts": nu.isoformat()}
                kandidaat = pk.get("kandidaat") or oud.get("kandidaat") or {}
                sinds = _tijd(kandidaat.get("ts"))
                laatst = _tijd(inv.get("laatst_gewijzigd"))
                dicht, reden = taken_dicht(rij, nu, self._lees_taakstatus)
                if len(inv["fotos"]) > vs.pakketcontract()["max_fotos"] or \
                        plan["totaal_bytes"] > vs.pakketcontract()["max_totaal_mb"] * 1_000_000:
                    self._nood(f"{mapnr}: pakket boven de foto- of totaalgrens; handmatig nakijken")
                elif bruikbaar(rij, inv, nu) and plan["transcripten"] and plan["fotos"] and dicht and \
                        sinds and laatst and nu - sinds >= timedelta(minutes=RUST_MIN) and \
                        nu - laatst >= timedelta(minutes=RUST_MIN):
                    uitslag = schrijf_pakket(plan, BS_INBOX, lambda p: bronnen.download(p, maximum=MAX_BESTAND),
                                            bronnen.download_kop, nu)
                    if uitslag["uitkomst"] == "grens":
                        pk["versiegrens"] = True
                        self._nood(f"{mapnr}: 5 versies in de inbox; nieuwe invoer wacht. Laat weten of oudere versies weg mogen; de agent ruimt niets op")
                    else:
                        pk["overdracht"] = uitslag
                        if uitslag["uitkomst"] == "klaargezet":
                            self.ag.log(onderwerp, "schrijf", f"pakket klaargezet voor het dashboard: {uitslag['pad']} "
                                        f"({uitslag['fotos']} foto's, {uitslag['bijlagen']} bijlagen); laden en verwerken doet een mens in barstenscheurenv2")
                        if uitslag["uitkomst"] in ("klaargezet", "al klaargezet"):
                            uitslag["fotos_iphone_bij_overdracht"] = len(inv["fotos_iphone"])
                            if uitslag["uitkomst"] == "klaargezet":
                                open_taken = _open_wachttaken(rij, self._lees_taakstatus)
                            else:
                                # Een al bestaand pakket is eerder geschreven; latere taken tellen niet mee.
                                grens = _tijd(uitslag.get("ts"))
                                open_taken = (_open_wachttaken(rij, self._lees_taakstatus, overdracht_ts=grens)
                                              if grens else [])
                            uitslag["open_wachttaak_bij_overdracht"] = bool(open_taken)
                            uitslag["open_wachttaak_reden_bij_overdracht"] = open_taken[0][1] if open_taken else None
                            uitslag["open_wachttaken_bij_overdracht"] = [uniek for uniek, _ in open_taken]
                        if plan["overgeslagen"]:
                            self._nood("overgeslagen uit pakket: " + ", ".join(os.path.basename(x["dropbox"] or "") for x in plan["overgeslagen"]))
        except Exception as e:  # noqa: BLE001
            self.ag.log(onderwerp, "fout", f"overdracht mislukt: {type(e).__name__}: {str(e)[:200]}")
            self._nood(f"{onderwerp}: overdracht mislukt ({type(e).__name__})")
        for k in ("overdracht", "kandidaat", "dossierkeuze", "bijlagen_na_overdracht", "versiegrens"):
            if k not in pk and k in oud:
                pk[k] = oud[k]
        overdracht = dict(pk.get("overdracht") or {})
        if overdracht:
            bij_overdracht = overdracht.get("fotos_iphone_bij_overdracht")
            if bij_overdracht is None:
                bij_overdracht = (0 if "zonder iPhone-foto's" in (overdracht.get("waarschuwing") or "")
                                  else len(inv["fotos_iphone"]))
            if "open_wachttaak_bij_overdracht" in overdracht:
                open_reden = (open_wachttaak(rij, self._lees_taakstatus,
                                             toegestane_taken=set(overdracht.get("open_wachttaken_bij_overdracht") or []))
                              if overdracht["open_wachttaak_bij_overdracht"] else None)
            else:
                grens = _tijd(overdracht.get("ts"))
                open_reden = open_wachttaak(rij, self._lees_taakstatus, overdracht_ts=grens) if grens else None
            if not bij_overdracht:
                waarschuwing = f"{mapnr}: pakket klaargezet zonder iPhone-foto's; controleer de fotolevering"
            elif open_reden:
                waarschuwing = (f"{mapnr}: pakket klaargezet na zes uur terwijl {open_reden}; "
                               "controleer of alle bronnen volledig zijn")
            else:
                waarschuwing = None
            if waarschuwing:
                overdracht["waarschuwing"] = waarschuwing
            elif overdracht.pop("waarschuwing", None):
                self.ag.log(onderwerp, "bevinding", "pakketwaarschuwing gewist: bronnen volledig en geen open wachttaken")
            pk["overdracht"] = overdracht
        vergelijk = lambda d: {k: v for k, v in d.items() if k != "gemeten_ts"}
        laatste_meting = _tijd(oud.get("gemeten_ts"))
        if vergelijk(pk) != vergelijk(oud) or not laatste_meting or nu - laatste_meting >= timedelta(hours=1):
            self._bewaar_pakket(rij, pk)
        waarschuwing = (pk.get("overdracht") or {}).get("waarschuwing")
        if waarschuwing_nood:
            if waarschuwing:
                waarschuwing_nood["tekst"] = waarschuwing
            else:
                self._noden.remove(waarschuwing_nood)
        elif waarschuwing:
            self._noden.insert(0, {"tekst": waarschuwing, "wie": "mehdi"})

    # ------------------------------------------------------------- ronde ---
    def ronde(self):
        self._noden = []
        self._mapcache = {}
        self._opbouwcache = {}
        self._inventaris = {}
        self._taakstatus = {}
        self.ag.hartslag("actief", taak="impulsen en opdrachten")
        try:
            try:
                n_i = self.impulsen()
            except Exception as e:  # noqa: BLE001
                n_i = 0
                self.ag.log("bak", "fout", f"impulsen niet leesbaar: {type(e).__name__}")
                self._nood(f"bak van het bord niet leesbaar ({type(e).__name__})")
            try:
                n_o = self.opdrachten()
            except Exception as e:  # noqa: BLE001
                n_o = 0
                self.ag.log("bak", "fout", f"opdrachten niet leesbaar: {type(e).__name__}")
                self._nood(f"bak van het bord niet leesbaar ({type(e).__name__})")
            gemeten = 0
            for rij in self.rijen():
                onderwerp = f"{rij.get('datum')} {rij.get('dossier') or rij.get('klant') or ''}".strip()
                try:
                    if self.s.get("submappen"):
                        self._ronde_dossiermap(rij)
                        gemeten += 1
                        continue
                    if not rij.get("dossiermap") and rij.get("impuls_voorbereiding_ts"):
                        pad, hoe = self.zoek_dossiermap(rij)
                        if pad:
                            self.voorbereiding(rij, onderwerp)
                            continue
                        self._noden.append({"tekst": f"{onderwerp}: dossiermap niet gevonden; zet het pad op de pagina Commandocentrum", "wie": "mehdi"})
                    elif rij.get("dossiermap") and not rij.get("bezoekmap"):
                        self.voorbereiding(rij, onderwerp)
                    if rij.get("bezoekmap") and not rij.get("proef_pad"):
                        pk = self.meet_pakket(rij)
                        oud = rij.get("pakket") or {}
                        if any(pk.get(k) != oud.get(k) for k in ("fotos", "opnames", "transcripten", "documenten", "bezoekmap")):
                            self._bewaar_pakket(rij, pk)
                            self.ag.log(onderwerp, "bevinding",
                                        f"pakket gemeten: {pk['fotos']} foto's, {pk['opnames']} opname(s), {pk['transcripten']} transcript(en), {pk['documenten']} document(en)")
                        gemeten += 1
                except Exception as e:  # noqa: BLE001
                    self.ag.log(onderwerp, "fout", f"ronde voor deze rij mislukt: {type(e).__name__}: {str(e)[:200]}")
                    self._nood(f"{onderwerp}: ronde voor deze rij mislukt ({type(e).__name__})")
            try:
                commit = subprocess.run(["git", "-C", os.path.dirname(HIER), "rev-parse", "--short", "HEAD"],
                                        capture_output=True, text=True, timeout=5, check=True).stdout.strip()
            except Exception:
                commit = "onbekend"
            self.ag.log_verstuur()
            self.ag.hartslag("klaar" if (n_i or n_o) else "waakt", taak=f"{n_i} impuls(en), {n_o} proef/proeven; {gemeten} pakket(ten) gemeten",
                             detail=f"impulsen van het Commandocentrum; proef alleen op de knop; commit {commit}", nood=self._noden[:6])
        except Exception as e:  # noqa: BLE001
            self.ag.log("", "fout", f"{type(e).__name__}: {str(e)[:300]}")
            self.ag.log_verstuur()
            self.ag.hartslag("fout", taak="ronde mislukt", detail=f"{type(e).__name__}: {str(e)[:120]}")
            raise

    def droog(self):
        """Alleen GET en Dropbox-lijsten; elke rij is één zelfstandige JSON-regel."""
        self._mapcache, self._opbouwcache, self._inventaris, self._taakstatus = {}, {}, {}, {}
        nu = datetime.now(timezone.utc)
        for rij in self.rijen():
            huidig = rij.get("dossiermap") or ""
            regel = {"id": rij.get("id"), "datum": rij.get("datum"), "dossier": rij.get("dossier"),
                     "adres": rij.get("adres"), "dossiermap": {"huidig": huidig, "gekozen": huidig},
                     "herkomst": "mehdi", "hoe": "", "nood": None, "tellingen": {}, "submappen": {},
                     "pakketnaam": None, "invoerhash": None, "taken_dicht": None,
                     "zou_overdragen": False, "reden": ""}
            try:
                if not self.s.get("submappen"):
                    regel["reden"] = "andere verslagsoort"
                else:
                    keuze = (rij.get("pakket") or {}).get("dossierkeuze") or {}
                    agentkeuze = keuze.get("door") == "agent" and keuze.get("pad") == huidig
                    regel["herkomst"] = "agent" if agentkeuze or not huidig else "mehdi"
                    if not huidig or (agentkeuze and not dossierkeuze_vast(rij)):
                        gekozen, hoe = self.zoek_dossiermap(rij)
                        regel["dossiermap"]["gekozen"] = gekozen
                        regel["hoe"] = hoe
                        regel["nood"] = self._zoeknood
                    else:
                        gekozen = huidig
                        if not vs.tkn_pad(gekozen, self.s["basismappen"][0]):
                            regel["nood"] = "dossiermap op bord heeft ongeldige vorm"
                            gekozen = None
                    if gekozen and not regel["nood"]:
                        test_rij = {**rij, "dossiermap": gekozen, "bezoekmap": gekozen}
                        pk = self.meet_pakket(test_rij)
                        regel["tellingen"] = {k: pk[k] for k in ("fotos", "fotos_iphone", "opnames", "transcripten", "documenten", "mails", "proef")}
                        regel["submappen"] = pk["submappen"]
                        inv = self._inventaris.get(rij["uniek"])
                        if inv:
                            plan = pakket_plan(test_rij, inv, self.s, vs.pakketcontract())
                            regel["pakketnaam"], regel["invoerhash"] = plan["naam"], plan["invoerhash"]
                            dicht, reden = taken_dicht(rij, nu, self._lees_taakstatus)
                            regel["taken_dicht"] = dicht
                            laatst = _tijd(inv["laatst_gewijzigd"])
                            pakket = rij.get("pakket") or {}
                            overdracht = pakket.get("overdracht") or {}
                            kandidaat_info = pakket.get("kandidaat") or {}
                            kandidaat = _tijd(kandidaat_info.get("ts")) if kandidaat_info.get("hash") == plan["invoerhash"] else None
                            if overdracht.get("invoerhash") == plan["invoerhash"]:
                                regel["reden"] = f"al overgedragen als {overdracht.get('naam') or plan['naam']}"
                            elif not dicht:
                                regel["reden"] = reden
                            elif not kandidaat:
                                regel["reden"] = "eerste ronde zet de kandidaat"
                            elif not laatst:
                                regel["reden"] = "laatste wijziging onbekend"
                            else:
                                rust_tot = max(kandidaat, laatst) + timedelta(minutes=RUST_MIN)
                                utc = lambda t: t.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
                                if nu < rust_tot:
                                    regel["reden"] = (f"rust loopt tot {utc(rust_tot)} (kandidaat sinds {utc(kandidaat)}, "
                                                      f"laatste wijziging {utc(laatst)})")
                                elif not bruikbaar(rij, inv, nu) or not plan["transcripten"] or not plan["fotos"]:
                                    regel["reden"] = "foto en transcript nog niet bruikbaar"
                                elif len(inv["fotos"]) > vs.pakketcontract()["max_fotos"] or \
                                        plan["totaal_bytes"] > vs.pakketcontract()["max_totaal_mb"] * 1_000_000:
                                    regel["reden"] = "pakket boven de foto- of totaalgrens"
                                else:
                                    regel["zou_overdragen"] = True
                                    regel["reden"] = "klaar voor overdracht"
            except Exception as e:  # noqa: BLE001
                regel["nood"] = f"droge controle mislukt ({type(e).__name__})"
            print(json.dumps(regel, ensure_ascii=False))

    def main(self):
        p = argparse.ArgumentParser()
        p.add_argument("--proef", type=int, metavar="ID", help="handmatig: proef voor rij <id> (kost tokens)")
        p.add_argument("--meet", action="store_true", help="alleen pakketten meten")
        p.add_argument("--droog", action="store_true", help="alleen lezen en JSON per rij tonen")
        a = p.parse_args()
        self._noden = []
        if a.droog:
            self.droog()
            return
        if a.proef:
            rij = self.rij(vid=a.proef)
            self.proef(rij, f"{rij['datum']} {rij.get('dossier') or rij.get('klant')}")
            self.ag.log_verstuur()
            return
        self.ronde()
