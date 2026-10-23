"""Verslag van bevindingen: het sjabloon van H-Architects voor een geschil of schadegeval.

Geen werfverslag. Een werfverslag volgt lopende werken op (doorlopende punten, veiligheid, volgend werfbezoek, akkoord binnen
vijf dagen) en krijgt een werfverslagnummer. Een verslag van bevindingen beschrijft op vraag van de bouwheer (of zijn
raadsman) de toestand na werken van derden: wat de architect zag, wat de bouwheer verklaarde en wat uit de stukken blijkt,
de technische beoordeling, de herstelling en een raming. Het krijgt geen werfverslagnummer en gaat niet naar de werfapp.
Eerste geval: 2309 Stockemstraat 15 Huldenberg (beslissing Mehdi 21-10-2026: "is niet een WV maar een verslag van bevinding").

Opbouw (identiek in Word en in markdown):
  kop          Projectnummer, titel "Verslag van bevindingen: <adres>", CONCEPT-regel, gegevenstabel
  1 opdracht   wie vroeg het verslag, waarvoor
  2 rol        eerdere betrokkenheid van het bureau bij het goed (feitelijk, met bron)
  3 werkwijze  plaatsbezoek, aanwezigen, beperkingen, herkomstcodes, ingeziene en ontbrekende stukken
  4 het goed   beschrijving van het goed en van de uitgevoerde werken
  5 stukken    offerte en facturen, met opmerkingen per post
  6 bevindingen per onderdeel: vaststelling, herkomst, beoordeling, gevolg, herstelling, foto's
  7 gemeld     door de bouwheer gemelde punten die niet ter plaatse werden vastgesteld (zonder oordeel)
  8 raming     herstelling per post, bedrag van-tot excl. btw
  9 voorbehoud vaste tekst
  10 ondertekening, foto's achteraan

Herkomst per bevinding: V = eigen vaststelling ter plaatse, B = verklaring van de bouwheer, S = blijkt uit stukken.
"""
import io
import json
import os
import re
import sys
from datetime import date

HIER = os.path.dirname(os.path.abspath(__file__))
if HIER not in sys.path:
    sys.path.insert(0, HIER)
import sjabloon_werfverslag as sj  # noqa: E402

SOORT = "verslag_bevindingen"
TITEL = "Verslag van bevindingen"
HERKOMST = {"V": "eigen vaststelling ter plaatse", "B": "verklaring van de bouwheer", "S": "blijkt uit de stukken"}

VOORBEHOUD = (
    "Dit verslag steunt op een visueel onderzoek tijdens het plaatsbezoek, zonder destructief onderzoek en zonder toegang tot "
    "het dak. Verborgen delen (onderdak, isolatie, houtstructuur, aansluitingen achter afwerkingen) werden niet onderzocht.\n"
    "Bij elke bevinding staat de herkomst. Verklaringen van de bouwheer en de inhoud van de stukken werden niet op hun "
    "juistheid gecontroleerd. Wie bij het bezoek niet aanwezig was, werd niet gehoord.\n"
    "Het verslag doet geen uitspraak over aansprakelijkheid, contractuele verplichtingen of de juridische kwalificatie van de "
    "gebreken. Dat is de taak van de raadsman van de bouwheer en, desgevallend, van een gerechtsdeskundige.\n"
    "De raming is een indicatieve orde van grootte, exclusief btw, en vervangt geen offerte van een herstelaannemer.\n"
    "Toestanden of herstellingen van na het plaatsbezoek zijn niet in dit verslag verwerkt.")


def leeg(nr="<nr>", adres="<adres van het goed>"):
    """Een leeg verslag met plaatshouders, als sjabloon voor Mehdi."""
    return {
        "dossier": nr, "adres": adres, "referentie": f"{nr}-VB<n>", "datum": "<JJJJ-MM-DD>", "opgemaakt": date.today().isoformat(),
        "opdrachtgever": "<naam bouwheer(s)>", "aanvrager": "<bouwheer / raadsman van de bouwheer>",
        "opdracht": ["<wie vroeg het verslag, wanneer en waarvoor>"],
        "rol": ["<eerdere betrokkenheid van H-Architects bij het goed, met bron (overeenkomst, vergunning, mails)>"],
        "werkwijze": ["<datum en duur van het plaatsbezoek, aanwezigen, weer, hulpmiddelen, wat niet bereikbaar was>"],
        "stukken": [{"stuk": "<offerte, facturen, foto's van de bouwheer, ...>", "datum": "", "opmerking": ""}],
        "ontbrekend": ["<stukken die het verslag zouden aanvullen maar niet werden voorgelegd>"],
        "goed": ["<beschrijving van het goed en van de uitgevoerde werken>"],
        "posten": [{"post": "<post uit de offerte of factuur>", "bron": "", "bedrag": "", "opmerking": ""}],
        "posten_toelichting": [],
        "onderdelen": [{"naam": "<onderdeel>", "bevindingen": [
            {"nr": "B1", "titel": "<titel>", "herkomst": "V", "vaststelling": "<wat er te zien was>", "beoordeling": "<technische beoordeling>",
             "gevolg": "<gevolg of schade>", "herstel": "<herstelling>", "fotos": [], "dias": ""}]}],
        "gemeld": [{"punt": "<door de bouwheer gemeld, niet ter plaatse vastgesteld>", "bron": ""}],
        "raming": [{"nr": "R1", "post": "<herstelling>", "bij": "B1", "van": 0, "tot": 0}],
        "raming_toelichting": [],
        "voorbehoud": VOORBEHOUD,
        "ondertekening": {"plaats": "Diest", "datum": "<dd maand jjjj>", "naam": "Mehdi Chegini", "hoedanigheid": "architect, zaakvoerder H-Architects BV"},
        "fotos_inleiding": "",
    }


def euro(x):
    """1234.5 -> '1.234,50'; tekst blijft tekst."""
    if isinstance(x, str):
        return x
    s = f"{x:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _bedrag(r):
    return isinstance(r.get("van"), (int, float)) and isinstance(r.get("tot"), (int, float))


def raming_totaal(v):
    """Som van de ramingen; een variant (in plaats van een andere post) en een p.m.-post tellen niet mee."""
    rijen = [r for r in v.get("raming", []) if not r.get("variant") and _bedrag(r)]
    return sum(r["van"] for r in rijen), sum(r["tot"] for r in rijen)


def raming_totalen(v):
    """Het totaal en, per variant die een post vervangt ('vervangt': 'R3a'), het totaal van dat scenario
    (driehoeksoverleg 22-10, K7: het scenario met nieuwe ramen mag niet verborgen blijven)."""
    van, tot = raming_totaal(v)
    uit = [("Totaal", van, tot)] if tot else []
    per_nr = {r.get("nr"): r for r in v.get("raming", [])}
    for r in v.get("raming", []):
        weg = per_nr.get(r.get("vervangt"))
        if r.get("variant") and _bedrag(r) and weg and _bedrag(weg):
            uit.append((f"Totaal met {r['nr']} in plaats van {weg['nr']}", van - weg["van"] + r["van"], tot - weg["tot"] + r["tot"]))
    return uit


def _ramingpost(x):
    if x.get("variant"):
        return x.get("post", "") + (f" (variant in plaats van {x['vervangt']})" if x.get("vervangt") else " (variant, niet in het totaal)")
    return x.get("post", "")


def _herkomst(code):
    """'V + B' -> 'eigen vaststelling ter plaatse, verklaring van de bouwheer'; vrije tekst (een bron van de schrijver) blijft staan."""
    code = (code or "").strip()
    if not re.fullmatch(r"[VBS](\s*[+,/]\s*[VBS])*", code):
        return code or "(na te kijken)"
    return ", ".join(HERKOMST[c] for c in re.findall(r"[VBS]", code))


def _ramingbedrag(x):
    if isinstance(x.get("van"), (int, float)) and isinstance(x.get("tot"), (int, float)):
        return f"{euro(x['van'])} tot {euro(x['tot'])}"
    return x.get("tekst") or "(raming door de architect in te vullen)"


def uit_werfproef(verslag):
    """De proef van de werfverslag schrijver (categorieën en punten) als verslag van bevindingen. Wat de schrijver niet
    weet (opdracht, rol van het bureau, stukken, beoordeling, herstelling) blijft '(in te vullen)' voor Mehdi."""
    nr = str(verslag["dossier"])
    bezoek = str(verslag.get("bezoek") or "")
    v = leeg(nr, verslag.get("adres", ""))
    open_ = "(in te vullen)"
    # het VB-nummer komt van de voorbereider (bordrijen per dossier); zonder VB-nummer niet gokken
    v.update({"referentie": f"{nr}-{bezoek}" if bezoek.startswith("VB") else f"{nr}-VB (nummer na te kijken)", "datum": verslag["datum"],
              "opgemaakt": verslag.get("opgemaakt") or date.today().isoformat(), "opdrachtgever": open_,
              "aanvrager": open_, "opdracht": [open_ + " wie vroeg het verslag, wanneer en waarvoor"],
              "rol": [open_ + " eerdere betrokkenheid van H-Architects bij dit goed, met bron"],
              "werkwijze": [x for x in (verslag.get("status_nl"), "Bronnen: " + verslag.get("bronnen_kort", "de bezoekmap") + ".") if x],
              "stukken": [], "ontbrekend": [], "goed": [open_], "posten": [], "gemeld": []})
    v["ondertekening"]["datum"] = open_
    i = 0
    v["onderdelen"] = []
    for cat in verslag.get("categorieen", []):
        bevindingen = []
        for p in cat.get("punten", []):
            if p.get("doorlopend"):
                continue   # vaste werfpunten horen niet in een verslag van bevindingen
            i += 1
            bevindingen.append({"nr": f"B{i}", "titel": p.get("titel", ""), "herkomst": p.get("bron") or "(na te kijken)",
                                "vaststelling": p.get("tekst_nl", ""), "beoordeling": open_, "herstel": open_,
                                "fotos": list(p.get("fotos", []))})
        if bevindingen:
            v["onderdelen"].append({"naam": cat["naam"], "bevindingen": bevindingen})
    v["raming"] = [{"nr": f"R{k}", "post": f"{x.get('onderdeel', '')}: {x.get('herstelling', '')}".strip(": "), "bij": "",
                    "tekst": x.get("raming") or ""} for k, x in enumerate(verslag.get("raming", []), 1)]
    return v


def als_werfverslag(v):
    """Alleen de fotoverwijzingen in de vorm van een werfverslag, voor het verzamelen van de foto's door de schrijver."""
    return {"categorieen": [{"naam": o["naam"], "punten": [{"fotos": list(b.get("fotos", []))} for b in o.get("bevindingen", [])]}
                            for o in v.get("onderdelen", [])]}


def _fotonummers(v, fotos=None):
    """Elke foto één nummer, in volgorde van eerste vermelding bij een bevinding."""
    nummer, lijst = {}, []
    for o in v.get("onderdelen", []):
        for b in o.get("bevindingen", []):
            for f in b.get("fotos", []):
                if (fotos is None or f in fotos) and f not in nummer:
                    nummer[f] = len(lijst) + 1
                    lijst.append((f, f"{b['nr']} {b['titel']}"))
    return nummer, lijst


def _fotoverwijzing(b, nummer):
    nrs = [nummer[f] for f in b.get("fotos", []) if f in nummer]
    return ("foto " if len(nrs) == 1 else "foto's ") + ", ".join(str(n) for n in nrs) if nrs else ""


# ------------------------------------------------------------------ markdown ---
def naar_markdown(v):
    r = [f"# {TITEL}: {v['dossier']} {v['adres']}", "",
         f"**CONCEPT** - opgemaakt op {v['opgemaakt']}. Na te kijken en aan te vullen door Mehdi Chegini vóór verzending.", "",
         "| | |", "|---|---|",
         f"| Referentie | {v.get('referentie', '')} |", f"| Opdrachtgever | {v.get('opdrachtgever', '')} |",
         f"| Adres van het goed | {v['adres']} |", f"| Datum plaatsbezoek | {sj.datum_nl(v['datum'])} |",
         f"| Verslagtype | {TITEL.lower()} (geen werfverslag) |",
         f"| Opgemaakt door | {sj.FIRMA['naam']}, {sj.FIRMA['adres']}, {sj.FIRMA['btw']} |", ""]
    nummer, _ = _fotonummers(v)
    k = 0

    def kop(t):
        nonlocal k
        k += 1
        r.extend([f"## {k}. {t}", ""])

    for titel, veld in (("Opdracht en voorwerp", "opdracht"), ("Rol van H-Architects in dit dossier", "rol"), ("Werkwijze", "werkwijze")):
        kop(titel)
        for alinea in v.get(veld, []):
            r.extend([alinea, ""])
        if veld == "werkwijze":
            r.extend(["Herkomst bij elke bevinding: " + "; ".join(f"{c} = {t}" for c, t in HERKOMST.items()) + ".", "",
                      "Ingeziene stukken:", "", "| Stuk | Datum | Opmerking |", "|---|---|---|"])
            r.extend(f"| {s.get('stuk', '')} | {s.get('datum', '')} | {s.get('opmerking', '')} |" for s in v.get("stukken", []))
            r.append("")
            if v.get("ontbrekend"):
                r.extend(["Niet voorgelegd:", ""] + [f"- {x}" for x in v["ontbrekend"]] + [""])
    kop("Het goed en de uitgevoerde werken")
    for alinea in v.get("goed", []):
        r.extend([alinea, ""])
    kop("Offerte, facturen en uitvoering")
    r.extend(["| Post | Bron | Bedrag excl. btw | Opmerking |", "|---|---|---|---|"])
    r.extend(f"| {p.get('post', '')} | {p.get('bron', '')} | {euro(p.get('bedrag', ''))} | {p.get('opmerking', '')} |" for p in v.get("posten", []))
    r.append("")
    for alinea in v.get("posten_toelichting", []):
        r.extend([alinea, ""])
    kop("Bevindingen")
    for o in v.get("onderdelen", []):
        r.extend([f"### {o['naam']}", ""])
        for b in o.get("bevindingen", []):
            ref = _fotoverwijzing(b, nummer)
            r.extend([f"**{b['nr']} {b['titel']}** - herkomst: {_herkomst(b.get('herkomst'))}" + (f" - {ref}" if ref else ""), ""])
            for label, veld in (("Vaststelling", "vaststelling"), ("Beoordeling", "beoordeling"), ("Gevolg", "gevolg"), ("Herstelling", "herstel")):
                if b.get(veld):
                    r.extend([f"*{label}:* {b[veld]}", ""])
            if b.get("dias"):
                r.extend([f"*Fotodossier van de bouwheer:* {b['dias']}", ""])
    kop("Door de bouwheer gemelde punten, niet ter plaatse vastgesteld")
    r.extend(f"- {g['punt']}" + (f" ({g['bron']})" if g.get("bron") else "") for g in v.get("gemeld", []))
    r.append("")
    kop("Herstelling en raming")
    r.extend(["| Nr | Herstelling | Bij | Raming excl. btw |", "|---|---|---|---|"])
    for x in v.get("raming", []):
        r.append(f"| {x.get('nr', '')} | {_ramingpost(x)} | {x.get('bij', '')} | {_ramingbedrag(x)} |")
    for label, van, tot in raming_totalen(v):
        r.append(f"| | {label} | | {euro(van)} tot {euro(tot)} euro |")
    r.append("")
    for alinea in v.get("raming_toelichting", []):
        r.extend([alinea, ""])
    kop("Voorbehoud")
    r.extend(v.get("voorbehoud", VOORBEHOUD).split("\n") + [""])
    kop("Ondertekening")
    o = v.get("ondertekening", {})
    r.extend([f"Opgemaakt te {o.get('plaats', '')} op {o.get('datum', '')}.", "", f"{o.get('naam', '')}, {o.get('hoedanigheid', '')}", ""])
    return "\n".join(r)


# ---------------------------------------------------------------------- docx ---
def naar_docx(v, fotos=None, master=sj.MASTER):
    """Word in de huisstijl van de contractmasters (zoals het werfverslag). fotos: dict bestandsnaam -> bytes."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm, Pt, RGBColor
    fotos = fotos or {}
    doc = Document(master) if os.path.exists(master) else Document()
    if os.path.exists(master):
        sj._leeg_lichaam(doc)
    BLAUW_H1, BLAUW_H2, GRIJS = RGBColor(0x36, 0x5F, 0x91), RGBColor(0x4F, 0x81, 0xBD), RGBColor(0x6B, 0x6B, 0x6B)
    sj._voettekst(doc, f"{v['dossier']} | {TITEL}")

    def alinea(tekst="", vet=False, cursief=False, grootte=None, kleur=None, na=None):
        p = doc.add_paragraph()
        if tekst:
            r = p.add_run(tekst); r.bold = vet; r.italic = cursief
            if grootte:
                r.font.size = Pt(grootte)
            if kleur is not None:
                r.font.color.rgb = kleur
        if na is not None:
            p.paragraph_format.space_after = Pt(na)
        return p

    def gelabeld(label, tekst, na=3):
        p = alinea(na=na)
        r = p.add_run(label + ": "); r.bold = True; r.font.color.rgb = BLAUW_H2
        p.add_run(tekst)
        return p

    def h1(t):
        p = alinea(t.upper(), vet=True, grootte=12, kleur=BLAUW_H1, na=6)
        p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True

    def h2(t):
        p = alinea(t, vet=True, kleur=BLAUW_H2, na=3)
        p.paragraph_format.space_before = Pt(9); p.paragraph_format.keep_with_next = True

    def tabel(koppen, rijen, breedtes=None):
        t = doc.add_table(rows=1 if koppen else 0, cols=len(koppen) if koppen else len(rijen[0]))
        t.style = "Table Grid"
        if koppen:
            for i, k in enumerate(koppen):
                c = t.rows[0].cells[i]; c.text = ""
                r = c.paragraphs[0].add_run(k); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = BLAUW_H1
        for rij in rijen:
            cellen = t.add_row().cells
            for i, val in enumerate(rij):
                cellen[i].text = ""
                r = cellen[i].paragraphs[0].add_run(str(val if val is not None else "")); r.font.size = Pt(9.5)
                if not koppen and i == 0:
                    r.bold = True
        if breedtes:
            for rij in t.rows:
                for i, b in enumerate(breedtes):
                    rij.cells[i].width = Cm(b)
        alinea(na=2)

    nummer, fotolijst = _fotonummers(v, fotos)
    alinea(f"Projectnummer: {v['dossier']}", vet=True, na=2)
    alinea(f"{TITEL}: {v['adres']}", vet=True, grootte=16, na=8)
    p = alinea(na=8)
    r = p.add_run("CONCEPT"); r.bold = True; r.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C)
    r = p.add_run(f" - opgemaakt op {v['opgemaakt']}. Na te kijken en aan te vullen door Mehdi Chegini vóór verzending.")
    r.italic = True; r.font.size = Pt(9.5)
    tabel(None, [["Referentie", v.get("referentie", "")], ["Opdrachtgever", v.get("opdrachtgever", "")], ["Adres van het goed", v["adres"]],
                 ["Datum plaatsbezoek", sj.datum_nl(v["datum"])], ["Opgemaakt door", f"{sj.FIRMA['naam']}, {sj.FIRMA['adres']}, {sj.FIRMA['btw']}, "
                                                                            f"{sj.FIRMA['mail']}, {sj.FIRMA['tel']}"]], [4.2, 12.2])
    k = 0

    def hoofdstuk(t):
        nonlocal k
        k += 1
        h1(f"{k}. {t}")

    hoofdstuk("Opdracht en voorwerp")
    for t in v.get("opdracht", []):
        alinea(t, na=4)
    hoofdstuk("Rol van H-Architects in dit dossier")
    for t in v.get("rol", []):
        alinea(t, na=4)
    hoofdstuk("Werkwijze")
    for t in v.get("werkwijze", []):
        alinea(t, na=4)
    alinea("Bij elke bevinding staat de herkomst: " + "; ".join(f"{c} = {t}" for c, t in HERKOMST.items()) + ".", na=6)
    h2("Ingeziene stukken")
    tabel(["Stuk", "Datum", "Opmerking"], [[s.get("stuk", ""), s.get("datum", ""), s.get("opmerking", "")] for s in v.get("stukken", [])], [7, 2.6, 6.8])
    if v.get("ontbrekend"):
        h2("Niet voorgelegd")
        for x in v["ontbrekend"]:
            alinea("- " + x, na=2)
    hoofdstuk("Het goed en de uitgevoerde werken")
    for t in v.get("goed", []):
        alinea(t, na=4)
    hoofdstuk("Offerte, facturen en uitvoering")
    tabel(["Post", "Bron", "Bedrag excl. btw", "Opmerking"],
          [[x.get("post", ""), x.get("bron", ""), euro(x.get("bedrag", "")), x.get("opmerking", "")] for x in v.get("posten", [])], [4.6, 2.4, 2.4, 7])
    for t in v.get("posten_toelichting", []):
        alinea(t, na=4)
    hoofdstuk("Bevindingen")
    for o in v.get("onderdelen", []):
        h2(o["naam"])
        for b in o.get("bevindingen", []):
            p = alinea(na=2); p.paragraph_format.keep_with_next = True
            r = p.add_run(f"{b['nr']}  {b['titel']}"); r.bold = True
            r = p.add_run(f"   herkomst: {_herkomst(b.get('herkomst'))}"); r.font.size = Pt(9); r.font.color.rgb = GRIJS
            ref = _fotoverwijzing(b, nummer)
            if ref:
                r = p.add_run(f"   ({ref})"); r.font.size = Pt(9); r.font.color.rgb = GRIJS
            for label, veld in (("Vaststelling", "vaststelling"), ("Beoordeling", "beoordeling"), ("Gevolg", "gevolg"), ("Herstelling", "herstel")):
                if b.get(veld):
                    gelabeld(label, b[veld])
            if b.get("dias"):
                gelabeld("Fotodossier van de bouwheer", b["dias"])
            alinea(na=4)
    hoofdstuk("Door de bouwheer gemelde punten, niet ter plaatse vastgesteld")
    alinea("Deze punten staan in de stukken van de bouwheer of werden mondeling gemeld. Ze werden tijdens het plaatsbezoek niet "
           "vastgesteld en worden hier zonder oordeel vermeld.", cursief=True, grootte=9.5, na=4)
    for g in v.get("gemeld", []):
        alinea("- " + g["punt"] + (f" ({g['bron']})" if g.get("bron") else ""), na=2)
    hoofdstuk("Herstelling en raming")
    rijen = [[x.get("nr", ""), _ramingpost(x), x.get("bij", ""), _ramingbedrag(x)] for x in v.get("raming", [])]
    for label, van, tot in raming_totalen(v):
        rijen.append(["", label, "", f"{euro(van)} tot {euro(tot)}"])
    tabel(["Nr", "Herstelling", "Bij", "Euro excl. btw"], rijen, [1.2, 9.8, 1.6, 3.8])
    for t in v.get("raming_toelichting", []):
        alinea(t, na=4)
    hoofdstuk("Voorbehoud")
    for t in v.get("voorbehoud", VOORBEHOUD).split("\n"):
        alinea(t, grootte=9.5, na=3)
    hoofdstuk("Ondertekening")
    o = v.get("ondertekening", {})
    alinea(f"Opgemaakt te {o.get('plaats', '')} op {o.get('datum', '')}.", na=18)
    alinea(o.get("naam", ""), vet=True, na=0)
    alinea(o.get("hoedanigheid", ""), na=6)
    if fotolijst:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        hoofdstuk("Foto's")
        alinea(v.get("fotos_inleiding") or f"{len(fotolijst)} foto's genomen door de architect tijdens het plaatsbezoek van {sj.datum_nl(v['datum'])}.",
               grootte=9.5, na=6)
        for i, (f, bij) in enumerate(fotolijst, 1):
            p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(0)
            try:
                p.add_run().add_picture(io.BytesIO(sj.beeld_rechtop(fotos[f])), height=Cm(10.6))
            except Exception as ex:  # noqa: BLE001
                p.add_run(f"(foto {f} niet leesbaar: {type(ex).__name__})")
            q = alinea(na=10)
            r = q.add_run(f"Foto {i}"); r.bold = True; r.font.size = Pt(9)
            tijd = f"{f[:2]}:{f[2:4]}" if f[:4].isdigit() else f
            r = q.add_run(f"   {tijd}   bij {bij}"); r.font.size = Pt(9); r.font.color.rgb = GRIJS
            if i % 2 == 0 and i < len(fotolijst):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def fotos_uit_map(map_):
    """Alle jpg/png uit een fotomap als dict naam -> bytes (HEIC kan niet in Word; de iCloud-wacht zet er een jpg naast)."""
    return {n: open(os.path.join(map_, n), "rb").read() for n in sorted(os.listdir(map_)) if n.lower().endswith((".jpg", ".jpeg", ".png"))}


if __name__ == "__main__":
    import argparse
    a = argparse.ArgumentParser(description="Verslag van bevindingen maken uit een gegevensbestand (JSON), of een leeg sjabloon.")
    a.add_argument("gegevens", nargs="?", help="JSON met de velden van leeg()")
    a.add_argument("--uit", required=True, help="pad van het Word-bestand; de markdown komt ernaast (.md)")
    a.add_argument("--fotos", help="map met de foto's (jpg/png) waarnaar de bevindingen verwijzen")
    a.add_argument("--leeg", action="store_true", help="leeg sjabloon met plaatshouders")
    arg = a.parse_args()
    v = leeg() if arg.leeg else json.load(open(arg.gegevens, encoding="utf8"))
    beelden = fotos_uit_map(arg.fotos) if arg.fotos else {}
    open(arg.uit, "wb").write(naar_docx(v, beelden))
    md = naar_markdown(v)
    open(os.path.splitext(arg.uit)[0] + ".md", "w", encoding="utf8").write(md)
    print(f"geschreven: {arg.uit} ({len(beelden)} foto's beschikbaar); open: {sj.telling_open(md)}")
