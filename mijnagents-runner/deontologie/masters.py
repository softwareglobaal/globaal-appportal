"""De harde controles van De Deontologiewacht op contracten (opdrachtkaart v0.3, stap 4 van de bouw).

Toetst de tekst van een contract aan de zes harde regels van het register: B1 tot B5 en D1. Elke controle zoekt naar
vaste formuleringen in de tekst en geeft een `Bevinding` met de plaats (het artikel van het contract) en een voorstel.
De kleur komt pas in `verslag.maak`, met de kleurregels van stap 3. Zo kan een controle hier nooit zelf rood of groen
maken: ze zegt alleen slaagt, faalt of onbekend.

Bewust streng en eenvoudig: wat de controle niet terugvindt, faalt. Een valse rode vlag kost Mehdi een blik; een
gemiste rode vlag staat in een contract van een klant. Twijfelt een controle (een uitsluiting die ze niet kan lezen), dan
zegt ze onbekend en wordt het oranje.

De masters zijn Jinja-sjablonen; de toets loopt op een gerenderde versie (de testdossiers van het contractsysteem), zodat
elke keuze in het sjabloon (budgetscenario, opdrachtomvang) apart getoetst wordt.

Gebruik:
    python3 -m deontologie.masters contract.docx [ander.txt ...]   (vanuit mijnagents-runner)
"""
import re
import sys
from pathlib import Path

from . import budget as G
from . import verslag as V

B = V.Bevinding
KOP = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+\S")
LOTEN = ("wind- en waterdicht", "technieken", "afwerking")

TAKEN = (  # de zeven taken van artikel 20 RBP2022, met de woorden waaraan ze in een contract te herkennen zijn
    ("de gegevens", r"verzamelen van (de )?gegevens|gegevens die nodig zijn"),
    ("het programma", r"behoeften van de opdrachtgever|programma|wenst volgende ruimtes"),
    ("schets en voorontwerp", r"schets|voorontwerp"),
    ("het administratief dossier", r"administratief dossier|aanvraag (van de |tot )?(omgevings|stedenbouwkundige )?vergunning|vergunningsaanvraag"),
    ("het dossier voor de bestelling", r"aanbesteding|meetsta|keuze van (de )?aannemer"),
    ("het uitvoeringsdossier met de controle der werken", r"uitvoeringsdossier.*|controle op de uitvoering"),
    ("de bijstand bij de oplevering en het nazicht der rekeningen", r"bijstand bij de oplevering|nazicht van de vorderingsstaten"),
)
UITSLUITING = (r"niet belast|niet inbegrepen|valt buiten|vallen buiten|omvat niet|omvat geen|geen deel uit|uitgesloten|"
               r"te beperken tot|beperkt tot|uitsluitend betrekking")
BEPERKT = r"te beperken tot|beperkt tot|uitsluitend betrekking|omvat geen werfopvolging"


def beperkte_opdracht(tekst):
    """Een voorstudie of regularisatie: de opdracht is uitdrukkelijk beperkt en er is geen uitvoeringsdossier. Of de
    regels over aannemer, begroting en programmawijziging daar gelden, is een keuze van Mehdi (opdrachtkaart,
    ontwerplezing punt 6: welke regel bij welke master hoort). Tot dan wordt zo een fout oranje, niet rood."""
    return bool(re.search(BEPERKT, tekst, re.I)) and not re.search(r"uitvoeringsdossier", tekst, re.I)


def _beperkt(b, tekst):
    if b.uitkomst == V.FAALT and beperkte_opdracht(tekst):
        return B(b.code, V.ONBEKEND, b.plaats, b.voorstel,
                 f"{b.toelichting}; beperkte opdracht, Mehdi beslist of deze regel hier geldt")
    return b


def secties(tekst):
    """De tekst per artikel: [(kop, [alinea's])]. Wat voor de eerste kop staat, heet 'aanhef'."""
    uit, kop, alineas = [], "aanhef", []
    for regel in (r.strip() for r in tekst.replace("\r", "").split("\n")):
        if not regel:
            continue
        if KOP.match(regel) and len(regel) < 120:
            uit.append((kop, alineas))
            kop, alineas = regel, []
        else:
            alineas.append(regel)
    uit.append((kop, alineas))
    return uit


def _klein(t):
    return re.sub(r"\s+", " ", t).casefold()


def _zoek(secs, patroon):
    """De eerste sectie (kop) waarin het patroon staat, in de kop of in een alinea; anders None."""
    p = re.compile(patroon, re.I)
    for kop, alineas in secs:
        if p.search(kop) or any(p.search(a) for a in alineas):
            return kop
    return None


def b1_schriftelijk(secs, tekst):
    partijen = re.search(r"\btussen\b", tekst, re.I) and re.search(r"de opdrachtgever", tekst, re.I) \
        and re.search(r"de architect", tekst, re.I)
    slot = re.search(r"(opgesteld|opgemaakt|gedaan) te [^,\n]+,? op \S+ \S+ \d{4}[^\n]*(onderteken|exempla)", tekst, re.I)
    if partijen and slot:
        return B("B1", V.SLAAGT, "slotbepaling")
    return B("B1", V.FAALT, "slotbepaling" if partijen else "aanhef",
             "Opgesteld te [plaats], op [datum], in twee exemplaren, ondertekend door de Opdrachtgever en de Architect "
             "voor de opdracht begint.",
             "geen partijen" if not partijen else "geen datum en ondertekening in het slot")


def b2_taken(secs, tekst):
    ontbreekt = []
    for naam, patroon in TAKEN:
        if _zoek(secs, patroon):
            continue
        ontbreekt.append(naam)
    if not ontbreekt:
        return B("B2", V.SLAAGT, "artikel over de taken van de architect")
    if re.search(UITSLUITING, tekst, re.I):
        # Er staat een uitsluiting, maar de controle kan niet lezen of die precies deze taken noemt.
        return B("B2", V.ONBEKEND, "artikel over de taken van de architect",
                 _b2_voorstel(ontbreekt), f"niet teruggevonden: {', '.join(ontbreekt)}; er staat wel een uitsluiting")
    return B("B2", V.FAALT, "artikel over de taken van de architect", _b2_voorstel(ontbreekt),
             f"niet teruggevonden: {', '.join(ontbreekt)}")


def _b2_voorstel(ontbreekt):
    return ("Zeg per taak van artikel 20 of de Architect ermee belast is, ook wat niet: \"De Architect is niet belast "
            f"met {ontbreekt[0]}.\" Ontbreekt nu: {', '.join(ontbreekt)}.")


def b3_controle(secs, tekst):
    k = _klein(tekst)
    if not re.search(r"(opmaak|opmaken|opstellen) van (het )?uitvoeringsdossier", k):
        return B("B3", V.SLAAGT, "", "", "geen uitvoeringsdossier in de opdracht")
    zonder = [lot for lot in LOTEN
              if re.search(rf"uitvoeringsdossier voor de {re.escape(lot)}", k)
              and not re.search(rf"controle op de uitvoering van de {re.escape(lot)}", k)]
    if not re.search(r"controle op de uitvoering", k):
        zonder = zonder or ["de werken"]
    if not zonder:
        return B("B3", V.SLAAGT, _zoek(secs, r"controle op de uitvoering") or "")
    if re.search(r"andere (ingeschreven )?architect.{0,80}controle", k):
        return B("B3", V.ONBEKEND, _zoek(secs, r"uitvoeringsdossier") or "", "",
                 "een andere architect wordt genoemd, nagaan of hij de controle krijgt")
    return B("B3", V.FAALT, _zoek(secs, r"uitvoeringsdossier") or "",
             f"Voeg toe: \"Controle op de uitvoering van de {zonder[0]}\", of noem de andere ingeschreven architect "
             "die de controle doet.", f"uitvoeringsdossier zonder controle: {', '.join(zonder)}")


def b4_aannemer(secs, tekst):
    patroon = r"bijstand bij de keuze van (de )?aannemer"
    kop = _zoek(secs, patroon)
    voorstel = "Neem op bij de taken: \"Bijstand bij de keuze van de aannemer(s) en de aanbesteding.\""
    if not kop:
        return B("B4", V.FAALT, "artikel over de taken van de architect", voorstel, "bijstand bij de keuze niet gevonden")
    alineas = dict(secs)[kop] if kop in dict(secs) else []
    zin = next((a for a in alineas if re.search(patroon, a, re.I)), kop)
    if re.search(r"\bniet\b|uitgesloten|behalve|\bgeen\b", zin, re.I):
        return B("B4", V.FAALT, kop, voorstel, "de bijstand staat er met een uitsluiting")
    return B("B4", V.SLAAGT, kop)


def b5_budget(secs, tekst):
    plaats = _zoek(secs, r"bouwbudget|begroting") or "artikel over het budget"
    if re.search(r"geen budgetverantwoordelijkheid|geen (bouw)?budget|zonder (bouw)?budget", tekst, re.I):
        return B("B5", V.FAALT, plaats, "Schrap de uitsluiting van de budgetverantwoordelijkheid en vermeld: \"Het "
                 "Bouwbudget wordt op [bedrag] euro vooropgesteld.\"", "de budgetverantwoordelijkheid wordt uitgesloten")
    programma = _zoek(secs, r"programma|wenst volgende ruimtes|ruimtes, functies")
    bedrag = None
    for kop, alineas in secs:
        if re.search(r"budget|begroting", kop, re.I):
            bedrag = bedrag or next((a for a in alineas if re.search(r"\d[\d.]*(,\d{2})?\s?(euro|eur\b|€)|€\s?\d", a, re.I)), None)
    if programma and bedrag and not re.search(r"(€|euro)\s*/\s*m|per (m²|m2|vierkante)", bedrag, re.I):
        return B("B5", V.SLAAGT, plaats)
    if programma and bedrag:
        return B("B5", V.ONBEKEND, plaats, "Vermeld naast de prijs per m² het totale Bouwbudget in euro.",
                 "alleen een prijs per m², geen totaal in euro")
    return B("B5", V.FAALT, plaats,
             "Vermeld bij ondertekening het programma en een begroting in euro: \"Het Bouwbudget wordt op [bedrag] euro "
             "vooropgesteld.\" Een raming die later volgt, komt in een Bijvoegsel.",
             "geen programma" if not programma else "geen begroting in euro bij het budget")


def d1_wijziging(secs, tekst):
    plaatsen, goed = [], None
    for kop, alineas in secs:
        blok = " ".join([kop] + alineas)
        if re.search(r"wijziging", blok, re.I) and re.search(r"bouwproject|programma", blok, re.I):
            plaatsen.append(kop)
            if re.search(r"bijvoegsel|addendum", blok, re.I) and re.search(r"financi|weerslag op het (bouw)?budget", blok, re.I):
                goed = kop
    if goed:
        return B("D1", V.SLAAGT, goed)
    return B("D1", V.FAALT, plaatsen[0] if plaatsen else "artikel over wijzigingen",
             "\"Elke wijziging van het programma wordt vastgelegd in een Bijvoegsel, ondertekend door beide Partijen, dat "
             "de financiele weerslag op het Bouwbudget en op het ereloon vermeldt.\"",
             "een wijziging van het programma vraagt geen addendum met de financiele weerslag")


CONTROLES = (b1_schriftelijk, b2_taken, b3_controle, b4_aannemer, b5_budget, d1_wijziging)


def toets_contract(tekst):
    """De bevindingen voor de harde regels; de andere regels blijven 'niet getoetst' (oranje) in het verslag."""
    secs = secties(tekst)
    uit = [c(secs, tekst) for c in CONTROLES]
    uit = [_beperkt(b, tekst) if b.code in ("B4", "B5", "D1") else b for b in uit]
    # Stap 7: budget tegen oppervlakte (D2, alleen cijfers zolang er geen referentie is) en de verzekering (D4).
    return uit + [G.d2_begroting(*G.uit_contract(tekst)), G.d4_verzekering(tekst)]


def lees(pad):
    pad = Path(pad)
    if pad.suffix.lower() == ".docx":
        from docx import Document
        d = Document(str(pad))
        delen = [p.text for p in d.paragraphs]
        for tabel in d.tables:
            for rij in tabel.rows:
                delen.append(" | ".join(c.text for c in rij.cells))
        return "\n".join(delen)
    return pad.read_text(encoding="utf-8")


if __name__ == "__main__":
    for pad in sys.argv[1:]:
        print(f"## {Path(pad).name}\n")
        print(V.maak("contract", toets_contract(lees(pad))).markdown())
        print()
