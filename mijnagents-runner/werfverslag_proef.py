"""Voorbereiden en proef maken voor De Werfverslaggever.

Twee stappen, naar het patroon van het contractsysteem (gegevens met herkomst, dan een proef):
  voorbereid(dossier, bezoek): leest alles in de bezoekmap (docx, pdf, pptx, md, txt), haalt er met
      claude-opus-5 de gegevens uit (bouwheer, aannemer, aanwezigen, doel, ...) met bron en zekerheid,
      plus de vaststellingen per onderdeel; zet ze op het bord.
  proef(dossier, bezoek): schrijft uit de gegevens en de bronteksten het concept-werfverslag
      (vaste opbouw E6 van de H-A vaste afspraken) als markdown en als Word, en zet beide in de
      bezoekmap in Dropbox als `<nr>-N werfverslag (concept).docx/.md`. Nooit overschrijven.

Regels: niets verzinnen (E8): wat niet in de bronnen staat, krijgt "(in te vullen)" of "(na te kijken)".
De sleutel ANTHROPIC_API_KEY komt uit ~/agents/.env; nooit in code of logboek.
"""
import io
import json
import os
import re
import sys
import zipfile
from datetime import date

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HIER, "koppelingen"))
import bord  # noqa: E402
import bronnen  # noqa: E402

MODEL = os.environ.get("WERFVERSLAG_MODEL", "claude-opus-5")
MAX_BRON_TEKENS = 60000


def _env(pad):
    pad = os.path.expanduser(pad)
    if not os.path.exists(pad):
        return
    for regel in open(pad):
        regel = regel.strip()
        if regel and not regel.startswith("#") and "=" in regel:
            k, v = regel.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip())


_env("~/agents/.env")
_env("~/appportal/.env")


# ------------------------------------------------------------ bronnen lezen ---
def docx_tekst(b):
    z = zipfile.ZipFile(io.BytesIO(b))
    x = z.read("word/document.xml").decode("utf8", "replace")
    x = re.sub(r"</w:p>", "\n", x)
    x = re.sub(r"<w:tab/>", "\t", x)
    return re.sub(r"[ \t]+", " ", re.sub(r"<[^>]+>", "", x)).strip()


def pptx_tekst(b):
    z = zipfile.ZipFile(io.BytesIO(b))
    dia = sorted((k for k in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", k)),
                 key=lambda k: int(re.search(r"(\d+)", k).group(1)))
    uit = []
    for k in dia:
        x = z.read(k).decode("utf8", "replace")
        t = " ".join(re.findall(r"<a:t>([^<]*)</a:t>", x))
        n = int(re.search(r"(\d+)", k).group(1))
        if t.strip():
            uit.append(f"dia {n}: {t.strip()}")
    media = sum(1 for k in z.namelist() if k.startswith("ppt/media/"))
    return f"({len(dia)} dia's, {media} afbeeldingen)\n" + "\n".join(uit)


def pdf_tekst(b, max_blz=12):
    try:
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(b))
        return "\n".join((p.extract_text() or "") for p in r.pages[:max_blz])
    except Exception as e:  # noqa: BLE001
        return f"(pdf niet leesbaar: {type(e).__name__})"


def lees_bezoekmap(bezoekmap):
    """Alle leesbare bestanden in de bezoekmap: [{naam, pad, tekst}], plus tellingen."""
    items = bronnen.lijst(bezoekmap) or []
    teksten, fotos, opnames = [], [], []
    for e in items:
        if e.get(".tag") != "file":
            continue
        pad, naam = e["path_display"], e["name"]
        laag = naam.lower()
        if laag.startswith(".") or "(concept)" in laag:
            continue
        try:
            if laag.endswith((".jpg", ".jpeg", ".heic", ".png")):
                fotos.append(naam)
            elif laag.endswith((".mp3", ".mp4", ".m4a", ".wav")):
                opnames.append(naam)
            elif laag.endswith(".docx"):
                teksten.append({"naam": naam, "pad": pad, "tekst": docx_tekst(bronnen.download(pad))})
            elif laag.endswith(".pptx"):
                teksten.append({"naam": naam, "pad": pad, "tekst": pptx_tekst(bronnen.download(pad))})
            elif laag.endswith(".pdf") and e.get("size", 0) < 6_000_000:
                teksten.append({"naam": naam, "pad": pad, "tekst": pdf_tekst(bronnen.download(pad))})
            elif laag.endswith((".md", ".txt")):
                teksten.append({"naam": naam, "pad": pad, "tekst": bronnen.download(pad).decode("utf8", "replace")})
        except Exception as ex:  # noqa: BLE001
            teksten.append({"naam": naam, "pad": pad, "tekst": f"(niet leesbaar: {type(ex).__name__})"})
    return teksten, fotos, opnames


def bronnenbundel(teksten, maximum=MAX_BRON_TEKENS):
    per = max(4000, maximum // max(1, len(teksten)))
    delen = []
    for t in teksten:
        tekst = t["tekst"]
        if len(tekst) > per:
            tekst = tekst[:per] + f"\n(... ingekort, {len(t['tekst'])} tekens in totaal)"
        delen.append(f"### BRON: {t['naam']}\n{tekst}")
    return "\n\n".join(delen)


# ------------------------------------------------------------- Claude ---
GEGEVENS_SCHEMA = {
    "name": "gegevens",
    "description": "De gegevens van het werfbezoek, elk met bron en zekerheid.",
    "input_schema": {
        "type": "object",
        "properties": {
            "gegevens": {"type": "array", "items": {"type": "object", "properties": {
                "veld": {"type": "string"},
                "waarde": {"type": "string"},
                "bron": {"type": "string", "description": "bestandsnaam en plaats waar het staat, of 'niet gevonden'"},
                "zekerheid": {"type": "string", "enum": ["zeker", "na te kijken", "ontbreekt"]},
            }, "required": ["veld", "waarde", "bron", "zekerheid"]}},
            "situatie": {"type": "string", "description": "de situatie op de werf in 5 tot 10 zinnen, alleen uit de bronnen"},
            "onderdelen": {"type": "array", "description": "vaststellingen per bouwonderdeel", "items": {"type": "object", "properties": {
                "onderdeel": {"type": "string"},
                "vaststellingen": {"type": "array", "items": {"type": "string"}},
                "bron": {"type": "string"},
            }, "required": ["onderdeel", "vaststellingen", "bron"]}},
            "acties": {"type": "array", "items": {"type": "object", "properties": {
                "wie": {"type": "string"}, "wat": {"type": "string"}, "tegen": {"type": "string"}}, "required": ["wie", "wat"]}},
            "ontbreekt": {"type": "array", "items": {"type": "string"}, "description": "wat de architect nog moet aanleveren of nakijken vóór het verslag af is"},
        },
        "required": ["gegevens", "situatie", "onderdelen", "acties", "ontbreekt"],
    },
}

VELDEN = ["bouwheer", "adres werf", "aannemer", "datum werfbezoek", "aanwezigen", "opdrachtgever van het verslag",
          "doel van het verslag", "contract H-Architects (opdracht)", "gebreken gemeld door de bouwheer",
          "schade en meerkosten gemeld door de bouwheer", "openstaande facturen", "juridische stand"]

SYSTEEM_VOORBEREID = """Je bent de voorbereider van werfverslagen van H-Architects (België). Je leest de bronnen van één
werfbezoek en haalt er de gegevens uit. Regel E8 van het kantoor: je verzint niets. Wat niet in de bronnen staat,
krijgt waarde "(in te vullen)" en zekerheid "ontbreekt". Wat je afleidt maar niet letterlijk leest, krijgt "na te
kijken". Elke waarde krijgt de bron (bestandsnaam, en waar in het bestand). Namen van externen: de firmanaam.
Datums als JJJJ-MM-DD. Schrijf in het Nederlands, zakelijk, zonder emoji. Vul minstens deze velden in: """ + ", ".join(VELDEN)

SYSTEEM_PROEF = """Je schrijft het concept-werfverslag van H-Architects (België) voor de klant, in markdown, in deze vaste
opbouw en precies deze koppen (regel E6):

# Werfverslag N - project <nr> <adres>
**CONCEPT** - opgemaakt op <datum> uit <bronnen>. Na te kijken en aan te vullen door Mehdi Chegini vóór verzending.
Tabel: Verslagnummer | Datum werfbezoek | Opgemaakt door (H-Architects BV, Mehdi Chegini) | Bronnen
## 1. Aanleiding en doel
## 2. Stand van de werken
## 3. Aanwezigen  (tabel Rol | Firma | Naam | Aanwezig)
## 4. Vaststellingen  (### per onderdeel, opsomming; verwijs naar de foto's/dia's waar de bron dat doet)
## 5. Raming van de herstelkosten  (tabel Onderdeel | Herstelling | Raming; bedragen alleen als ze in de bronnen staan,
   anders "(raming door de architect in te vullen)")
## 6. Actiepunten  (tabel Wie | Wat | Tegen)
## 7. Volgend werfbezoek
## 8. Algemene voorwaarden  (vaste tekst: veiligheidscoördinatie; tienjarige aansprakelijkheid wet Peeters-Borsus van
   31 mei 2017; "Zonder tegenbericht per e-mail binnen de vijf werkdagen wordt aangenomen dat alle partijen akkoord
   gaan met dit verslag. Vragen of opmerkingen kunnen per e-mail worden bezorgd aan H-Architects.")
## 9. Voor akkoord  (tabel Partij | Naam | Handtekening | Datum: Opdrachtgever, Architect H-Architects bv - Mehdi Chegini,
   Aannemer; en de regel "Opgemaakt en verzonden door de architect op: <datum>")

Regels: niets verzinnen; wat onzeker is krijgt "(na te kijken)", wat ontbreekt "(in te vullen)". Geen juridische
conclusies (wie aansprakelijk is): beschrijf vaststellingen en verwijs voor de rest naar de raadsman van de bouwheer.
Hoofdstukken zonder inhoud krijgen "geen". Nederlands, zakelijk, geen emoji, geen kastlijntjes (gebruik een gewoon
koppelteken). Alleen de markdown, geen inleiding of nawoord."""


def _client():
    from anthropic import Anthropic
    return Anthropic()


def vraag_gegevens(kop, bundel):
    resp = _client().messages.create(model=MODEL, max_tokens=6000, system=SYSTEEM_VOORBEREID,
                                     messages=[{"role": "user", "content": kop + "\n\n" + bundel}],
                                     tools=[GEGEVENS_SCHEMA], tool_choice={"type": "tool", "name": "gegevens"})
    for blok in resp.content:
        if getattr(blok, "type", "") == "tool_use":
            return blok.input, resp.usage
    return {}, resp.usage


def vraag_proef(kop, gegevens, bundel):
    user = kop + "\n\n## GEGEVENS (uit de voorbereiding, met bron en zekerheid)\n" + json.dumps(gegevens, ensure_ascii=False, indent=1) + "\n\n## BRONNEN\n" + bundel
    resp = _client().messages.create(model=MODEL, max_tokens=9000, system=SYSTEEM_PROEF,
                                     messages=[{"role": "user", "content": user}])
    return "".join(getattr(b, "text", "") for b in resp.content).strip(), resp.usage


# ---------------------------------------------------------- markdown -> docx ---
def md_naar_docx(md):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    stijl = doc.styles["Normal"]
    stijl.font.name = "Calibri"
    stijl.font.size = Pt(10.5)
    regels = md.splitlines()
    i = 0

    def inline(par, tekst):
        for stuk in re.split(r"(\*\*[^*]+\*\*)", tekst):
            if stuk.startswith("**") and stuk.endswith("**"):
                par.add_run(stuk[2:-2]).bold = True
            elif stuk:
                par.add_run(stuk)

    while i < len(regels):
        r = regels[i].rstrip()
        if not r.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", r)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 3))
            i += 1
            continue
        if r.lstrip().startswith("|"):
            rijen = []
            while i < len(regels) and regels[i].lstrip().startswith("|"):
                cellen = [c.strip() for c in regels[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cellen if c) or not any(cellen):
                    rijen.append(cellen)
                i += 1
            if rijen:
                kol = max(len(x) for x in rijen)
                tabel = doc.add_table(rows=0, cols=kol)
                tabel.style = "Table Grid"
                for n, rij in enumerate(rijen):
                    cellen = tabel.add_row().cells
                    for k in range(kol):
                        p = cellen[k].paragraphs[0]
                        inline(p, rij[k] if k < len(rij) else "")
                        if n == 0:
                            for run in p.runs:
                                run.bold = True
            continue
        if re.match(r"^\s*[-*]\s+", r):
            inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", r))
            i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", r):
            inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", r))
            i += 1
            continue
        # gewone alinea: aaneengesloten regels samenvoegen
        alinea = [r.strip()]
        i += 1
        while i < len(regels) and regels[i].strip() and not re.match(r"^(#|\||\s*[-*]\s|\s*\d+\.\s)", regels[i]):
            alinea.append(regels[i].strip())
            i += 1
        inline(doc.add_paragraph(), " ".join(alinea))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- stappen ---
def _rij(dossier, volgnr):
    r = bord.call(f"/api/werfbezoek?dossier={dossier}&volgnr={volgnr}")
    rijen = r.get("rijen") or []
    if not rijen:
        raise RuntimeError(f"bezoek {dossier}-{volgnr} staat niet op het bord; draai eerst een verificatieronde")
    return rijen[0]


def voorbereid(ag, dossier, volgnr):
    rij = _rij(dossier, volgnr)
    ag.log(f"{dossier}-{volgnr}", "bron", f"bezoekmap lezen: {rij['bezoekmap']}")
    teksten, fotos, opnames = lees_bezoekmap(rij["bezoekmap"])
    ag.log(f"{dossier}-{volgnr}", "bron", f"{len(teksten)} tekstbron(nen), {len(fotos)} foto's, {len(opnames)} opname(s) gelezen",
           [t["naam"] + f" ({len(t['tekst'])} tekens)" for t in teksten])
    kop = (f"Dossier {dossier}, {rij.get('adres','')}. Werfbezoek {volgnr} op {rij['datum']}. "
           f"Bezoekmap: {rij['bezoekmap']}. Foto's in de map: {len(fotos)}; opnames: {len(opnames)}.")
    uit, usage = vraag_gegevens(kop, bronnenbundel(teksten))
    uit["bronbestanden"] = [t["naam"] for t in teksten]
    uit["fotos_in_map"] = fotos
    uit["opnames_in_map"] = opnames
    ag.log(f"{dossier}-{volgnr}", "bevinding", f"{len(uit.get('gegevens', []))} gegevens, {len(uit.get('onderdelen', []))} onderdelen, "
           f"{len(uit.get('ontbreekt', []))} open punten ({usage.input_tokens}+{usage.output_tokens} tokens)")
    bord.call("/api/werfbezoek", {"rijen": [{"dossier": dossier, "datum": rij["datum"], "bezoekmap": rij["bezoekmap"],
                                             "volgnr": volgnr, "gegevens": uit, "_alleen": ["gegevens"]}]})
    return uit


def proef(ag, dossier, volgnr):
    rij = _rij(dossier, volgnr)
    gegevens = rij.get("gegevens") or {}
    if not gegevens.get("gegevens"):
        gegevens = voorbereid(ag, dossier, volgnr)
    teksten, fotos, opnames = lees_bezoekmap(rij["bezoekmap"])
    vandaag = date.today().isoformat()
    kop = (f"Dossier {dossier}, {rij.get('adres','')}. Werfbezoek {volgnr} op {rij['datum']}; verslagnummer {dossier}-{volgnr}. "
           f"Datum van opmaak: {vandaag}. Foto's in de map: {len(fotos)}; opnames: {len(opnames)}.")
    md, usage = vraag_proef(kop, gegevens, bronnenbundel(teksten))
    if not md.startswith("#"):
        md = f"# Werfverslag {volgnr} - project {dossier} {rij.get('adres','')}\n\n" + md
    naam = f"{dossier}-{volgnr} werfverslag (concept)"
    pad_md = bronnen.upload(f"{rij['bezoekmap']}/{naam}.md", md.encode("utf8"))
    pad_docx = bronnen.upload(f"{rij['bezoekmap']}/{naam}.docx", md_naar_docx(md))
    ag.log(f"{dossier}-{volgnr}", "schrijf", f"proef gezet: {os.path.basename(pad_docx)} en {os.path.basename(pad_md)} in de bezoekmap "
           f"({usage.input_tokens}+{usage.output_tokens} tokens)")
    bord.call("/api/werfbezoek", {"rijen": [{"dossier": dossier, "datum": rij["datum"], "bezoekmap": rij["bezoekmap"], "volgnr": volgnr,
                                             "verslag_md": md, "proef_pad": pad_docx, "proef_ts": vandaag, "_alleen": ["verslag_md", "proef_pad", "proef_ts"]}]})
    return pad_docx, md
